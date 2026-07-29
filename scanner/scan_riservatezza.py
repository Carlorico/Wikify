#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Scanner deterministico di riservatezza - Metodo A (assessment Archivio Smart)

Percorre un albero di cartelle, estrae il testo dai file (docx, xlsx, pdf,
txt e simili), applica il dizionario dei pattern e produce un report XLSX:
  - Segnalazioni:        una riga per ogni corrispondenza trovata
  - Riepilogo file:      una riga per file, con conteggi per severita'
  - Riepilogo categorie: totali per categoria di regola
  - Non analizzati:      file non estraibili (legacy, scansioni, errori)
  - Registro segregazione (bozza): precompilato per il triage del
    responsabile tecnico (Qualifica = "Da valutare")

Tutta l'elaborazione avviene in locale. Nessun contenuto lascia la macchina.

Uso:
    python scan_riservatezza.py --root <cartella_archivio> [opzioni]

Opzioni:
    --root       cartella radice da scansionare (obbligatoria)
    --dizionario percorso del dizionario YAML (default: dizionario_pattern.yaml
                 nella cartella dello script)
    --out        file di report (default: report_scansione_<data>.xlsx)
    --contesto   caratteri di contesto attorno al match (default: 80)

Dipendenze: pyyaml, openpyxl. Facoltativa: pypdf (per i PDF).
    pip install pyyaml openpyxl pypdf
"""

import argparse
import datetime
import os
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

try:
    import yaml
except ImportError:
    sys.exit("Manca pyyaml. Installare con: pip install pyyaml")

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    sys.exit("Manca openpyxl. Installare con: pip install openpyxl")

try:
    from pypdf import PdfReader
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

# ---------------------------------------------------------------------------
# Estrazione testo per formato
# Ogni estrattore restituisce una lista di (posizione, testo).
# ---------------------------------------------------------------------------

TEXT_EXT = {".txt", ".csv", ".md", ".log", ".ini", ".cfg", ".json", ".xml"}
LEGACY_EXT = {".doc", ".xls", ".ppt"}
NS_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def extract_txt(path):
    for enc in ("utf-8", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as f:
                lines = f.readlines()
            return [(f"riga {i+1}", line.rstrip("\n")) for i, line in enumerate(lines) if line.strip()]
        except UnicodeDecodeError:
            continue
    raise ValueError("codifica non riconosciuta")


def extract_docx(path):
    """Estrae i paragrafi da un .docx leggendo direttamente l'XML (nessuna dipendenza)."""
    out = []
    with zipfile.ZipFile(path) as z:
        targets = [n for n in z.namelist()
                   if n == "word/document.xml" or re.match(r"word/(header|footer)\d*\.xml", n)]
        for name in targets:
            root = ET.fromstring(z.read(name))
            for i, p in enumerate(root.iter(f"{NS_W}p")):
                text = "".join(t.text or "" for t in p.iter(f"{NS_W}t"))
                if text.strip():
                    origine = "corpo" if name == "word/document.xml" else os.path.basename(name)
                    out.append((f"{origine}, par. {i+1}", text))
    return out


def extract_xlsx(path):
    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and str(cell.value).strip():
                    out.append((f"foglio '{ws.title}', cella {cell.coordinate}", str(cell.value)))
    wb.close()
    return out


def extract_pdf(path):
    if not HAS_PDF:
        raise ValueError("pypdf non installato (pip install pypdf)")
    reader = PdfReader(path)
    out = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        for j, line in enumerate(text.splitlines()):
            if line.strip():
                out.append((f"pag. {i+1}, riga {j+1}", line))
    if not out:
        raise ValueError("nessun testo estraibile (possibile PDF scansionato: valutare OCR)")
    return out


def extract(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXT_EXT:
        return extract_txt(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext in (".xlsx", ".xlsm"):
        return extract_xlsx(path)
    if ext == ".pdf":
        return extract_pdf(path)
    if ext in LEGACY_EXT:
        raise ValueError("formato legacy: convertire in formato recente o esaminare manualmente")
    raise ValueError("formato non gestito dallo scanner")


# ---------------------------------------------------------------------------
# Dizionario e scansione
# ---------------------------------------------------------------------------

def load_rules(path):
    with open(path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    rules = []
    for r in data.get("regole", []):
        try:
            rx = re.compile(r["pattern"], re.IGNORECASE)
        except re.error as e:
            print(f"  ATTENZIONE: regola {r.get('id','?')} ignorata, regex non valida ({e})")
            continue
        rules.append({**r, "rx": rx})
    if not rules:
        sys.exit("Nessuna regola valida nel dizionario.")
    return rules


def scan_file(path, rules, ctx_len):
    """Restituisce (segnalazioni, n_unita_testo). Solleva ValueError se non estraibile."""
    units = extract(path)
    findings = []
    for pos, text in units:
        for rule in rules:
            for m in rule["rx"].finditer(text):
                a, b = m.start(), m.end()
                ctx = text[max(0, a - ctx_len): b + ctx_len].strip()
                findings.append({
                    "regola": rule["id"], "categoria": rule["categoria"],
                    "severita": rule["severita"], "descrizione": rule["descrizione"],
                    "posizione": pos, "match": m.group(0), "contesto": ctx,
                })
    return findings, len(units)


def walk(root):
    skip_prefix = ("~$", ".")
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if not d.startswith(".")]
        for name in sorted(filenames):
            if name.startswith(skip_prefix):
                continue
            yield os.path.join(dirpath, name)


# ---------------------------------------------------------------------------
# Report XLSX
# ---------------------------------------------------------------------------

SEV_ORDER = {"alta": 0, "media": 1, "bassa": 2}
SEV_FILL = {
    "alta": PatternFill("solid", start_color="F4CCCC"),
    "media": PatternFill("solid", start_color="FCE5CD"),
    "bassa": PatternFill("solid", start_color="FFF2CC"),
}


def autosize(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def write_report(out_path, root, all_findings, file_stats, not_analyzed, rules):
    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    # --- Segnalazioni ---
    ws = wb.active
    ws.title = "Segnalazioni"
    header = ["File", "Posizione", "Regola", "Categoria", "Severità",
              "Testo intercettato", "Contesto"]
    ws.append(header)
    for c in ws[1]:
        c.font = bold
    ordered = sorted(all_findings, key=lambda f: (SEV_ORDER.get(f["severita"], 9), f["file"]))
    for f in ordered:
        ws.append([f["file"], f["posizione"], f["regola"], f["categoria"],
                   f["severita"], f["match"], f["contesto"]])
        ws.cell(row=ws.max_row, column=5).fill = SEV_FILL.get(f["severita"], PatternFill())
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws, {"A": 55, "B": 22, "C": 9, "D": 18, "E": 9, "F": 30, "G": 70})

    # --- Riepilogo file ---
    ws = wb.create_sheet("Riepilogo file")
    ws.append(["File", "Formato", "Unità di testo", "Segnalazioni",
               "Alta", "Media", "Bassa", "Esito"])
    for c in ws[1]:
        c.font = bold
    for st in sorted(file_stats, key=lambda s: -s["tot"]):
        ws.append([st["file"], st["ext"], st["units"], st["tot"],
                   st["alta"], st["media"], st["bassa"],
                   "DA TRIAGE" if st["tot"] else "nessun pattern (campionare a controllo)"])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autosize(ws, {"A": 55, "B": 9, "C": 13, "D": 13, "E": 7, "F": 7, "G": 7, "H": 34})

    # --- Riepilogo categorie ---
    ws = wb.create_sheet("Riepilogo categorie")
    ws.append(["Categoria", "Regole attive", "Segnalazioni", "File coinvolti"])
    for c in ws[1]:
        c.font = bold
    cats = {}
    for r in rules:
        cats.setdefault(r["categoria"], {"rules": 0, "n": 0, "files": set()})["rules"] += 1
    for f in all_findings:
        cats[f["categoria"]]["n"] += 1
        cats[f["categoria"]]["files"].add(f["file"])
    for cat, d in sorted(cats.items(), key=lambda kv: -kv[1]["n"]):
        ws.append([cat, d["rules"], d["n"], len(d["files"])])
    autosize(ws, {"A": 24, "B": 13, "C": 13, "D": 14})

    # --- Non analizzati ---
    ws = wb.create_sheet("Non analizzati")
    ws.append(["File", "Motivo"])
    for c in ws[1]:
        c.font = bold
    for path, reason in not_analyzed:
        ws.append([path, reason])
    autosize(ws, {"A": 55, "B": 60})

    # --- Registro segregazione (bozza) ---
    ws = wb.create_sheet("Registro segregazione (bozza)")
    ws.append(["File / sezione", "Pattern rilevato", "Qualifica",
               "Motivazione", "Vincolo", "Data", "Validatore"])
    for c in ws[1]:
        c.font = bold
    seen = set()
    for f in ordered:
        key = (f["file"], f["regola"])
        if key in seen:
            continue
        seen.add(key)
        ws.append([f"{f['file']} ({f['posizione']})", f["regola"], "Da valutare",
                   "", "", "", ""])
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"Riservato,Condivisibile,Da valutare"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    if ws.max_row > 1:
        dv.add(f"C2:C{ws.max_row}")
    ws.freeze_panes = "A2"
    autosize(ws, {"A": 60, "B": 16, "C": 15, "D": 40, "E": 25, "F": 12, "G": 16})

    # --- Parametri scansione ---
    ws = wb.create_sheet("Parametri scansione")
    for k, v in [("Radice scansionata", root),
                 ("Data e ora", datetime.datetime.now().strftime("%d/%m/%Y %H:%M")),
                 ("Regole caricate", len(rules)),
                 ("File analizzati", len(file_stats)),
                 ("File non analizzati", len(not_analyzed)),
                 ("Segnalazioni totali", len(all_findings)),
                 ("Nota", "Elaborazione interamente locale. Nessun contenuto trasmesso all'esterno.")]:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = bold
    autosize(ws, {"A": 24, "B": 80})

    wb.save(out_path)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description="Scanner deterministico di riservatezza (Metodo A)")
    ap.add_argument("--root", required=True, help="cartella radice da scansionare")
    ap.add_argument("--dizionario", default=None, help="dizionario YAML dei pattern")
    ap.add_argument("--out", default=None, help="file XLSX di report")
    ap.add_argument("--contesto", type=int, default=80, help="caratteri di contesto")
    args = ap.parse_args()

    root = os.path.abspath(args.root)
    if not os.path.isdir(root):
        sys.exit(f"Cartella non trovata: {root}")

    dict_path = args.dizionario or os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                                "dizionario_pattern.yaml")
    rules = load_rules(dict_path)
    print(f"Dizionario: {dict_path} ({len(rules)} regole)")
    print(f"Scansione di: {root}")

    all_findings, file_stats, not_analyzed = [], [], []
    for path in walk(root):
        rel = os.path.relpath(path, root)
        try:
            findings, units = scan_file(path, rules, args.contesto)
        except ValueError as e:
            not_analyzed.append((rel, str(e)))
            continue
        except Exception as e:
            not_analyzed.append((rel, f"errore di lettura: {e}"))
            continue
        for f in findings:
            f["file"] = rel
        all_findings.extend(findings)
        sev = {"alta": 0, "media": 0, "bassa": 0}
        for f in findings:
            sev[f["severita"]] = sev.get(f["severita"], 0) + 1
        file_stats.append({"file": rel, "ext": os.path.splitext(path)[1].lower(),
                           "units": units, "tot": len(findings), **sev})

    out = args.out or f"report_scansione_{datetime.date.today():%Y%m%d}.xlsx"
    write_report(out, root, all_findings, file_stats, not_analyzed, rules)

    print(f"\nFile analizzati: {len(file_stats)}  |  non analizzati: {len(not_analyzed)}")
    print(f"Segnalazioni: {len(all_findings)} "
          f"(alta: {sum(1 for f in all_findings if f['severita']=='alta')}, "
          f"media: {sum(1 for f in all_findings if f['severita']=='media')}, "
          f"bassa: {sum(1 for f in all_findings if f['severita']=='bassa')})")
    print(f"Report: {out}")


if __name__ == "__main__":
    main()
