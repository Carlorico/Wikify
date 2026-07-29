# -*- coding: utf-8 -*-
"""
Modulo Validazione AI: export del perimetro condivisibile, import dei
lotti di bozze dell'agente, coda di revisione con vista affiancata,
note di conoscenza tacita, metriche ed export.

Convenzione sugli endpoint: i nomi iniziano per "perimetro", "coda",
"metriche" o "impostazioni", cosi' la sidebar evidenzia la sotto-voce
corretta.
"""

import html as html_mod
import io
import json
import os

from flask import (Blueprint, Response, flash, redirect, render_template,
                   request, session, url_for)
from markupsafe import Markup

from core import analisi as ana
from core import db
from core import inventario as inv
from core import validazione as val
from core.report import sanifica_cella

bp = Blueprint("validazione", __name__)

STATI_CODA = ("da_fare", "validate", "tutte")

ESITI_LABEL = {"confermata": "Confermata", "corretta": "Corretta",
               "caso_nuovo": "Caso nuovo"}


def _utente_sessione():
    return session.get("utente_nome") or ""


def _radice(conn):
    return db.get_impostazione(conn, inv.CHIAVE_RADICE, "")


def _scansioni_completate(conn):
    return conn.execute(
        "SELECT * FROM scansioni WHERE stato = 'completata' "
        "ORDER BY id DESC").fetchall()


def _scansione_scelta(conn):
    """La scansione indicata in querystring, oppure l'ultima completata."""
    scansioni = _scansioni_completate(conn)
    voluta = request.args.get("scansione", "")
    scelta = None
    if voluta.isdigit():
        for s in scansioni:
            if s["id"] == int(voluta):
                scelta = s
                break
    if scelta is None and scansioni:
        scelta = scansioni[0]
    return scansioni, scelta


def _lotto_o_none(conn, lid):
    return conn.execute("SELECT * FROM lotti_validazione WHERE id = ?",
                        (lid,)).fetchone()


def _lotti_con_avanzamento(conn):
    return conn.execute(
        "SELECT l.*, "
        "(SELECT COUNT(*) FROM proposte p JOIN validazioni v "
        " ON v.proposta_id = p.id WHERE p.lotto_id = l.id) AS n_validate "
        "FROM lotti_validazione l ORDER BY l.id DESC").fetchall()


def _evidenzia(testo, citazione):
    """Testo HTML con la citazione avvolta in <mark> (se presente)."""
    esc = html_mod.escape(testo or "")
    cit = html_mod.escape(citazione or "")
    if cit:
        indice = esc.lower().find(cit.lower())
        if indice >= 0:
            esc = (esc[:indice] + "<mark>" + esc[indice:indice + len(cit)]
                   + "</mark>" + esc[indice + len(cit):])
    return Markup(esc)


# ---------------------------------------------------------------------------
# Perimetro e bozze
# ---------------------------------------------------------------------------

@bp.route("/")
def perimetro_home():
    return redirect(url_for("validazione.perimetro"))


@bp.route("/perimetro")
def perimetro():
    conn = db.get_connection()
    scansioni, scelta = _scansione_scelta(conn)
    dati_perimetro = val.perimetro_scansione(conn, scelta) if scelta else None
    lotti = _lotti_con_avanzamento(conn)
    dettagli_lotti = {}
    for lotto in lotti:
        try:
            dettagli_lotti[lotto["id"]] = json.loads(
                lotto["dettaglio_scarti_json"] or "{}")
        except ValueError:
            dettagli_lotti[lotto["id"]] = {}
    conn.close()
    return render_template(
        "validazione_perimetro.html", scansioni=scansioni, scelta=scelta,
        perimetro=dati_perimetro, lotti=lotti,
        motivi_label=val.MOTIVI_SCARTO_LABEL)


@bp.route("/perimetro/export.json")
def perimetro_export():
    conn = db.get_connection()
    _scansioni, scelta = _scansione_scelta(conn)
    if scelta is None:
        conn.close()
        flash("Nessuna scansione completata: eseguire prima la scansione "
              "e il triage.", "errore")
        return redirect(url_for("validazione.perimetro"))
    dati = val.perimetro_scansione(conn, scelta)
    conn.close()
    corpo = json.dumps(dati, ensure_ascii=False, indent=2)
    return Response(corpo, mimetype="application/json",
                    headers={"Content-Disposition":
                             "attachment; filename=perimetro_condivisibile.json"})


DOCS_AGENTE = {
    "istruzioni": ("istruzioni_agente.md", "istruzioni_agente.md"),
    "formato": ("formato_bozze.md", "formato_bozze.md"),
}


@bp.route("/perimetro/documenti/<nome>")
def perimetro_documenti(nome):
    """Download dei documenti per l'agente (istruzioni e formato bozze)."""
    voce = DOCS_AGENTE.get(nome)
    if voce is None:
        flash("Documento non disponibile.", "errore")
        return redirect(url_for("validazione.perimetro"))
    percorso = os.path.join(os.path.dirname(os.path.dirname(
        os.path.abspath(__file__))), "docs_agente", voce[0])
    if not os.path.isfile(percorso):
        flash("Documento non trovato nell'installazione.", "errore")
        return redirect(url_for("validazione.perimetro"))
    with open(percorso, "r", encoding="utf-8") as f:
        corpo = f.read()
    return Response(corpo, mimetype="text/markdown; charset=utf-8",
                    headers={"Content-Disposition":
                             "attachment; filename=%s" % voce[1]})


@bp.route("/perimetro/importa", methods=["POST"])
def perimetro_importa():
    files = [f for f in request.files.getlist("bozze") if f and f.filename]
    if not files:
        flash("Selezionare almeno un file JSON di bozza da importare.",
              "errore")
        return redirect(url_for("validazione.perimetro"))
    conn = db.get_connection()
    ultimo_lotto = None
    for f in files:
        try:
            dati = json.loads(f.read().decode("utf-8-sig"))
        except (ValueError, UnicodeDecodeError):
            flash("File rifiutato (%s): non è un JSON leggibile."
                  % f.filename, "errore")
            continue
        lotto_id, esito = val.importa_bozza(conn, f.filename, dati)
        if lotto_id is None:
            flash("File rifiutato (%s): %s." % (f.filename, esito), "errore")
            continue
        ultimo_lotto = lotto_id
        messaggio = ("Lotto %d creato da %s: %d proposte accettate, "
                     "%d scartate" % (lotto_id, f.filename,
                                      esito["n_accettate"],
                                      esito["n_scartate"]))
        if esito["n_troncate"]:
            messaggio += (", %d citazioni oltre %d caratteri troncate"
                          % (esito["n_troncate"], val.LIMITE_CITAZIONE))
        flash(messaggio + ".", "ok")
    conn.close()
    if ultimo_lotto is not None and len(files) == 1:
        return redirect(url_for("validazione.perimetro_lotto",
                                lid=ultimo_lotto))
    return redirect(url_for("validazione.perimetro"))


@bp.route("/perimetro/lotti/<int:lid>")
def perimetro_lotto(lid):
    conn = db.get_connection()
    lotto = _lotto_o_none(conn, lid)
    if lotto is None:
        conn.close()
        flash("Lotto non trovato.", "errore")
        return redirect(url_for("validazione.perimetro"))
    try:
        dettaglio = json.loads(lotto["dettaglio_scarti_json"] or "{}")
    except ValueError:
        dettaglio = {}
    try:
        note_agente = json.loads(lotto["note_agente_json"] or "[]")
    except ValueError:
        note_agente = []
    proposte = conn.execute(
        "SELECT p.*, v.esito, v.validatore, v.data AS data_validazione "
        "FROM proposte p LEFT JOIN validazioni v ON v.proposta_id = p.id "
        "WHERE p.lotto_id = ? ORDER BY p.confidenza ASC, p.id",
        (lid,)).fetchall()
    n_validate = sum(1 for p in proposte if p["esito"])
    conn.close()
    return render_template(
        "validazione_lotto.html", lotto=lotto, dettaglio=dettaglio,
        note_agente=note_agente, proposte=proposte, n_validate=n_validate,
        motivi_label=val.MOTIVI_SCARTO_LABEL, esiti_label=ESITI_LABEL)


# ---------------------------------------------------------------------------
# Coda di revisione
# ---------------------------------------------------------------------------

@bp.route("/coda")
def coda():
    conn = db.get_connection()
    cartelle = val.cartelle_in_coda(conn)
    conn.close()
    return render_template("validazione_coda.html", cartelle=cartelle)


@bp.route("/coda/revisione")
def coda_revisione():
    cartella = (request.args.get("cartella") or "").strip()
    stato = request.args.get("stato", "da_fare")
    if stato not in STATI_CODA:
        stato = "da_fare"
    if not cartella:
        return redirect(url_for("validazione.coda"))
    conn = db.get_connection()
    radice = _radice(conn)
    tutte = val.proposte_cartella(conn, cartella)
    conteggi = {"tutte": len(tutte),
                "validate": sum(1 for p in tutte if p["esito"]),
                "da_fare": sum(1 for p in tutte if not p["esito"])}
    if stato == "da_fare":
        mostrate = [p for p in tutte if not p["esito"]]
    elif stato == "validate":
        mostrate = [p for p in tutte if p["esito"]]
    else:
        mostrate = list(tutte)

    schede = []
    for p in mostrate:
        estratto = val.estratto_documento(
            radice, cartella, p["file"], p["evidenza_posizione"])
        unita_html = None
        if estratto is not None:
            unita_html = [
                {"pos": pos,
                 "html": (_evidenzia(testo, p["evidenza_citazione"])
                          if e_match else Markup(html_mod.escape(testo))),
                 "match": e_match}
                for pos, testo, e_match in estratto["unita"]]
        schede.append({
            "p": p, "unita": unita_html,
            "avvertenza": val.avvertenza_scanner(
                p["campo"], p["valore_proposto"], p["esito"] or ""),
            "tassonomia": val.TASSONOMIE.get(p["campo"]),
        })
    lotto_note = conn.execute(
        "SELECT MAX(lotto_id) FROM proposte WHERE cartella_progetto = ?",
        (cartella,)).fetchone()[0]
    note = conn.execute(
        "SELECT * FROM note_conoscenza_tacita WHERE cartella_progetto = ? "
        "ORDER BY id DESC", (cartella,)).fetchall()
    conn.close()
    return render_template(
        "validazione_revisione.html", cartella=cartella, stato=stato,
        conteggi=conteggi, schede=schede, note=note, lotto_note=lotto_note,
        esiti_label=ESITI_LABEL, avvertenza_testo=val.AVVERTENZA_SCANNER)


def _redirect_revisione(cartella):
    return redirect(url_for("validazione.coda_revisione", cartella=cartella,
                            stato=request.form.get("filtro_stato")
                            or "da_fare"))


@bp.route("/coda/proposte/<int:pid>/valida", methods=["POST"])
def coda_valida(pid):
    conn = db.get_connection()
    proposta = conn.execute("SELECT * FROM proposte WHERE id = ?",
                            (pid,)).fetchone()
    if proposta is None:
        conn.close()
        flash("Proposta non trovata.", "errore")
        return redirect(url_for("validazione.coda"))
    cartella = proposta["cartella_progetto"]
    esito = request.form.get("esito") or ""
    nota = (request.form.get("nota") or "").strip()
    try:
        secondi = max(0, int(float(request.form.get("secondi") or 0)))
    except ValueError:
        secondi = 0
    valore_corretto = ""
    errore = None
    if esito not in val.ESITI:
        errore = "Esito di validazione non valido."
    elif esito == "corretta":
        valore_corretto = (request.form.get("valore_corretto") or "").strip()
        tassonomia = val.TASSONOMIE.get(proposta["campo"])
        if not valore_corretto:
            errore = "Indicare il valore corretto."
        elif tassonomia and valore_corretto not in tassonomia:
            errore = ("Il valore corretto deve appartenere alla tassonomia "
                      "del campo %s." % proposta["campo"])
        elif valore_corretto == proposta["valore_proposto"]:
            errore = ("Il valore corretto coincide con quello proposto: "
                      "usare Conferma.")
    elif esito == "caso_nuovo":
        valore_corretto = (request.form.get("valore_nuovo") or "").strip()
        if not valore_corretto:
            errore = "Descrivere il caso nuovo nel campo valore."
    if errore:
        conn.close()
        flash(errore, "errore")
        return _redirect_revisione(cartella)
    avvertenza = val.salva_validazione(
        conn, proposta, esito, valore_corretto, nota, _utente_sessione(),
        secondi)
    conn.close()
    messaggio = "Proposta %s (%s, %s)." % (
        ESITI_LABEL[esito].lower(), proposta["campo"], proposta["file"])
    flash(messaggio, "ok")
    if avvertenza:
        flash(val.AVVERTENZA_SCANNER + ".", "attenzione")
    return _redirect_revisione(cartella)


@bp.route("/coda/note", methods=["POST"])
def coda_nota():
    cartella = (request.form.get("cartella") or "").strip()
    testo = (request.form.get("testo") or "").strip()
    lotto_id = request.form.get("lotto_id") or ""
    if not cartella or not testo:
        flash("Indicare il testo della nota.", "errore")
        return redirect(url_for("validazione.coda"))
    conn = db.get_connection()
    conn.execute(
        "INSERT INTO note_conoscenza_tacita (lotto_id, cartella_progetto, "
        "testo, autore, data) VALUES (?, ?, ?, ?, ?)",
        (int(lotto_id) if lotto_id.isdigit() else None, cartella, testo,
         _utente_sessione(), val._adesso()))
    conn.commit()
    conn.close()
    flash("Nota di conoscenza tacita registrata.", "ok")
    return _redirect_revisione(cartella)


# ---------------------------------------------------------------------------
# Metriche
# ---------------------------------------------------------------------------

def _lotto_filtro(conn):
    voluto = request.args.get("lotto", "")
    if voluto.isdigit():
        lotto = _lotto_o_none(conn, int(voluto))
        if lotto is not None:
            return lotto
    return None


def _sessione_filtro(conn):
    voluto = request.args.get("sessione", "")
    if voluto.isdigit():
        riga = conn.execute(
            "SELECT * FROM sessioni_analisi WHERE id = ?",
            (int(voluto),)).fetchone()
        if riga is not None:
            return riga
    return None


@bp.route("/metriche")
def metriche():
    conn = db.get_connection()
    lotto = _lotto_filtro(conn)
    sessione_sel = _sessione_filtro(conn)
    complessive = val.metriche(conn)
    del_lotto = val.metriche(conn, lotto_id=lotto["id"]) if lotto else None
    del_sessione = (val.metriche(conn, sessione_id=sessione_sel["id"])
                   if sessione_sel else None)
    lotti = _lotti_con_avanzamento(conn)
    sessioni = ana.elenco_sessioni(conn)
    kpi_sessioni = [ana.kpi_sessione(conn, s) for s in sessioni]
    conn.close()
    return render_template(
        "validazione_metriche.html", complessive=complessive,
        del_lotto=del_lotto, lotto=lotto, del_sessione=del_sessione,
        sessione_sel=sessione_sel, lotti=lotti, sessioni=sessioni,
        kpi_sessioni=kpi_sessioni, campi=val.CAMPI,
        motivi_label=val.MOTIVI_SCARTO_LABEL)


# ---------------------------------------------------------------------------
# Impostazioni analisi (connettore MCP, progettazione_moduli.md Modulo 4)
# ---------------------------------------------------------------------------

@bp.route("/impostazioni", methods=["GET", "POST"])
def impostazioni_analisi():
    conn = db.get_connection()
    if request.method == "POST":
        ok, errori = ana.valida_e_salva(conn, request.form)
        if ok:
            flash("Impostazioni di analisi salvate.", "ok")
        else:
            for errore in errori:
                flash(errore, "errore")
    impostazioni = ana.leggi_impostazioni(conn)
    conn.close()
    return render_template(
        "validazione_impostazioni.html", impostazioni=impostazioni,
        modelli_suggeriti=ana.MODELLI_SUGGERITI)


def _scrivi_metriche(ws, bold, titolo, m):
    ws.append([titolo])
    ws[ws.max_row][0].font = bold
    ws.append(["Accuratezza per campo", "N. validate", "Confermate",
               "Accuratezza %"])
    for c in ws[ws.max_row]:
        c.font = bold
    for campo in val.CAMPI:
        dati = m["per_campo"][campo]
        ws.append([campo, dati["n"], dati["confermate"],
                   dati["accuratezza"] if dati["accuratezza"] is not None
                   else ""])
    ws.append([])
    ws.append(["Riservatezza", "N. validate", "Confermate", "Accuratezza %"])
    for c in ws[ws.max_row]:
        c.font = bold
    for nome, etichetta in (("puri", "Documento intero (puri)"),
                            ("misti", "Sezioni (misti)")):
        dati = m["riservatezza"][nome]
        ws.append([etichetta, dati["n"], dati["confermate"],
                   dati["accuratezza"] if dati["accuratezza"] is not None
                   else ""])
    ws.append([])
    ws.append(["Fascia di confidenza", "N. validate", "Confermate",
               "Accuratezza %"])
    for c in ws[ws.max_row]:
        c.font = bold
    for f in m["fasce"]:
        ws.append([f["fascia"], f["n"], f["confermate"],
                   f["accuratezza"] if f["accuratezza"] is not None else ""])
    ws.append([])
    ws.append(["Tempo mediano per proposta (s)",
               m["tempo_mediano_proposta"] if m["tempo_mediano_proposta"]
               is not None else ""])
    ws.append(["Tempo mediano per cartella (s)",
               m["tempo_mediano_cartella"] if m["tempo_mediano_cartella"]
               is not None else ""])
    ws.append(["Proposte validate", m["n_validate"]])
    ws.append([])


@bp.route("/metriche/export.xlsx")
def metriche_export_xlsx():
    import openpyxl
    from openpyxl.styles import Font
    conn = db.get_connection()
    lotto = _lotto_filtro(conn)
    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Metriche"
    _scrivi_metriche(ws, bold, "Metriche complessive", val.metriche(conn))
    if lotto is not None:
        _scrivi_metriche(ws, bold, "Metriche del lotto %d (%s)"
                         % (lotto["id"], lotto["nome_file_origine"]),
                         val.metriche(conn, lotto["id"]))

    ws = wb.create_sheet("Per lotto")
    ws.append(["Lotto", "File di origine", "Campo", "N. validate",
               "Confermate", "Accuratezza %"])
    for c in ws[1]:
        c.font = bold
    for riga_lotto in conn.execute(
            "SELECT id, nome_file_origine FROM lotti_validazione ORDER BY id"):
        m = val.metriche(conn, riga_lotto["id"])
        for campo in val.CAMPI:
            dati = m["per_campo"][campo]
            if not dati["n"]:
                continue
            ws.append([riga_lotto["id"],
                       sanifica_cella(riga_lotto["nome_file_origine"]),
                       campo, dati["n"], dati["confermate"],
                       dati["accuratezza"]])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Casi nuovi")
    ws.append(["Cartella progetto", "File", "Campo", "Valore proposto",
               "Caso nuovo (valore libero)", "Nota", "Validatore", "Data"])
    for c in ws[1]:
        c.font = bold
    m = val.metriche(conn, lotto["id"] if lotto else None)
    for r in m["casi_nuovi"]:
        ws.append([sanifica_cella(v) for v in (
            r["cartella_progetto"], r["file"], r["campo"],
            r["valore_proposto"], r["valore_corretto"], r["nota_val"],
            r["validatore"], r["data_val"])])
    ws.freeze_panes = "A2"
    conn.close()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(buf.read(),
                    mimetype="application/vnd.openxmlformats-officedocument"
                             ".spreadsheetml.sheet",
                    headers={"Content-Disposition":
                             "attachment; filename=metriche_validazione.xlsx"})


@bp.route("/metriche/correzioni.json")
def metriche_export_json():
    conn = db.get_connection()
    lotto = _lotto_filtro(conn)
    correzioni = val.correzioni_export(conn, lotto["id"] if lotto else None)
    conn.close()
    corpo = json.dumps({"formato": "wikify-correzioni/1.0",
                        "data_export": val._adesso(),
                        "correzioni": correzioni},
                       ensure_ascii=False, indent=2)
    return Response(corpo, mimetype="application/json",
                    headers={"Content-Disposition":
                             "attachment; filename=correzioni_validazione.json"})
