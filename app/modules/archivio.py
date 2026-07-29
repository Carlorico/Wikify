# -*- coding: utf-8 -*-
"""
Modulo Archivio logico (Modulo 5): regole di tipologia con builder guidato
e prova live, consolidamento con report, esplorazione entita' -> tipologia
-> documenti, completezza per entita', export JSON e XLSX. La logica di
dominio vive in core/archivio.py.
"""

import io
import json

from flask import (Blueprint, Response, flash, redirect, render_template,
                   request, url_for)

from core import archivio as arc
from core import db
from core.report import sanifica_cella

bp = Blueprint("archivio", __name__)


# ---------------------------------------------------------------------------
# Regole di tipologia
# ---------------------------------------------------------------------------

@bp.route("/")
def regole():
    """Elenco delle regole di tipologia, ordinate per priorità."""
    conn = db.get_connection()
    righe = arc.regole_tutte(conn)
    conn.close()
    return render_template("archivio_regole.html", regole=righe)


def _regola_o_none(conn, rid):
    return conn.execute("SELECT * FROM regole_tipologia WHERE id = ?",
                        (rid,)).fetchone()


def _form_regola(conn, regola):
    """GET/POST comuni di creazione e modifica di una regola di tipologia.

    Come nel builder del dizionario e del catalogo: il pulsante "Prova sui
    nomi reali" (azione=prova) calcola il conteggio dei documenti
    intercettati; solo dopo la prova compare "Salva"."""
    prova, errore = None, None
    if request.method == "POST":
        azione = request.form.get("azione") or "prova"
        nome = (request.form.get("nome") or "").strip()
        tipologia = request.form.get("tipologia") or ""
        criterio, errore = arc.criterio_da_form(request.form)
        if errore is None and not nome:
            errore = "Indicare il nome della regola (es. \"Manuali di installazione\")."
        if errore is None and tipologia not in arc.TASSONOMIA_TIPOLOGIA:
            errore = "Scegliere la tipologia da assegnare dall'elenco."
        if errore is None:
            prova = arc.prova_criterio(conn, criterio)
            if azione == "salva":
                if regola is None:
                    arc.crea_regola(conn, nome, tipologia, criterio)
                else:
                    arc.aggiorna_regola(conn, regola["id"], nome, tipologia,
                                       criterio)
                conn.close()
                flash("Regola di tipologia salvata.", "ok")
                return redirect(url_for("archivio.regole"))
        if errore:
            flash(errore, "errore")
    parametri, tipologia_sel = {}, ""
    if regola is not None and request.method == "GET":
        criterio_salvato = json.loads(regola["criterio_json"] or "{}")
        parametri = {"campo": criterio_salvato.get("campo", "nome_file"),
                    "modo": criterio_salvato.get("modo", "contiene"),
                    "valori": "\n".join(criterio_salvato.get("valori") or [])}
        tipologia_sel = regola["tipologia"]
    conn.close()
    return render_template(
        "archivio_regola_form.html", regola=regola, prova=prova,
        campi=arc.CAMPI_OSSERVATI, modi=arc.MODI,
        tipologie=arc.TASSONOMIA_TIPOLOGIA, parametri=parametri,
        tipologia_sel=tipologia_sel,
        form=request.form if request.method == "POST" else None)


@bp.route("/regole/nuova", methods=["GET", "POST"])
def regola_nuova():
    return _form_regola(db.get_connection(), None)


@bp.route("/regole/<int:rid>/modifica", methods=["GET", "POST"])
def regola_modifica(rid):
    conn = db.get_connection()
    regola = _regola_o_none(conn, rid)
    if regola is None:
        conn.close()
        flash("Regola di tipologia non trovata.", "errore")
        return redirect(url_for("archivio.regole"))
    return _form_regola(conn, regola)


@bp.route("/regole/<int:rid>/toggle", methods=["POST"])
def regola_toggle(rid):
    conn = db.get_connection()
    arc.toggle_regola(conn, rid)
    conn.close()
    return redirect(url_for("archivio.regole"))


@bp.route("/regole/<int:rid>/elimina", methods=["POST"])
def regola_elimina(rid):
    conn = db.get_connection()
    arc.elimina_regola(conn, rid)
    conn.close()
    flash("Regola di tipologia eliminata.", "ok")
    return redirect(url_for("archivio.regole"))


@bp.route("/regole/<int:rid>/sposta", methods=["POST"])
def regola_sposta(rid):
    direzione = request.form.get("direzione") or "su"
    conn = db.get_connection()
    arc.sposta_priorita(conn, rid, direzione)
    conn.close()
    return redirect(url_for("archivio.regole"))


# ---------------------------------------------------------------------------
# Consolidamento
# ---------------------------------------------------------------------------

@bp.route("/consolidamento", methods=["GET", "POST"])
def consolidamento():
    conn = db.get_connection()
    tipi = conn.execute(
        "SELECT id, nome FROM tipi_entita ORDER BY nome").fetchall()
    tipo_sel = request.values.get("tipo_entita", "")
    if not tipo_sel:
        tipo_sel = db.get_impostazione(conn, arc.IMPOSTAZIONE_TIPO_ENTITA, "")
    report = None
    if request.method == "POST":
        tipo_id = int(tipo_sel) if str(tipo_sel).isdigit() else None
        if tipo_id:
            db.set_impostazione(conn, arc.IMPOSTAZIONE_TIPO_ENTITA,
                                str(tipo_id))
        report = arc.ricostruisci(conn, tipo_id)
        flash("Consolidamento eseguito.", "ok")
    indicatori = arc.indicatori_copertura(conn)
    conn.close()
    return render_template("archivio_consolidamento.html", tipi=tipi,
                           tipo_sel=tipo_sel, report=report,
                           indicatori=indicatori)


# ---------------------------------------------------------------------------
# Esplora archivio
# ---------------------------------------------------------------------------

@bp.route("/esplora")
def esplora():
    conn = db.get_connection()
    entita_arg = request.args.get("entita", "")
    tipologia = request.args.get("tipologia", "")
    riservatezza = request.args.get("riservatezza", "")
    stato = request.args.get("stato", "")
    cerca = bool(request.args.get("cerca"))

    senza_entita = entita_arg == "senza"
    eid = int(entita_arg) if entita_arg.isdigit() else None

    righe_entita, n_senza = arc.riepilogo_entita(conn)
    tipologie_riepilogo = None
    if entita_arg and not tipologia:
        tipologie_riepilogo = arc.riepilogo_tipologie(conn, eid, senza_entita)

    documenti = None
    mostra_documenti = bool(tipologia or riservatezza or stato or cerca)
    if mostra_documenti:
        documenti = arc.documenti_filtrati(
            conn, eid, tipologia or None, riservatezza or None,
            stato or None, senza_entita)
    conn.close()
    return render_template(
        "archivio_esplora.html", righe_entita=righe_entita, n_senza=n_senza,
        entita_arg=entita_arg, tipologia=tipologia,
        riservatezza=riservatezza, stato=stato,
        tipologie_riepilogo=tipologie_riepilogo, documenti=documenti,
        tipologie=arc.TASSONOMIA_TIPOLOGIA + [arc.TIPOLOGIA_NON_CLASSIFICATO])


@bp.route("/documenti/<int:did>")
def documento_scheda(did):
    conn = db.get_connection()
    doc = conn.execute(
        "SELECT d.*, e.chiave AS entita_chiave, e.id AS entita_id_reale, "
        "t.nome AS entita_tipo_nome, r.nome AS regola_nome "
        "FROM documenti d LEFT JOIN entita e ON e.id = d.entita_id "
        "LEFT JOIN tipi_entita t ON t.id = e.tipo_id "
        "LEFT JOIN regole_tipologia r ON r.id = d.tipologia_regola_id "
        "WHERE d.id = ?", (did,)).fetchone()
    conn.close()
    if doc is None:
        flash("Documento non trovato.", "errore")
        return redirect(url_for("archivio.esplora"))
    return render_template("archivio_documento.html", doc=doc,
                           tipologie=arc.TASSONOMIA_TIPOLOGIA)


@bp.route("/documenti/<int:did>/correggi", methods=["POST"])
def documento_correggi(did):
    tipologia = request.form.get("tipologia") or ""
    conn = db.get_connection()
    doc = conn.execute("SELECT id FROM documenti WHERE id = ?",
                       (did,)).fetchone()
    valori_ammessi = arc.TASSONOMIA_TIPOLOGIA + [arc.TIPOLOGIA_NON_CLASSIFICATO]
    if doc is None:
        conn.close()
        flash("Documento non trovato.", "errore")
        return redirect(url_for("archivio.esplora"))
    if tipologia not in valori_ammessi:
        conn.close()
        flash("Scegliere una tipologia valida dall'elenco.", "errore")
        return redirect(url_for("archivio.documento_scheda", did=did))
    arc.correggi_tipologia_manuale(conn, did, tipologia)
    conn.close()
    flash("Tipologia corretta manualmente: la correzione non verrà "
          "sovrascritta dai prossimi consolidamenti.", "ok")
    return redirect(url_for("archivio.documento_scheda", did=did))


# ---------------------------------------------------------------------------
# Completezza per entità
# ---------------------------------------------------------------------------

@bp.route("/completezza")
def completezza():
    conn = db.get_connection()
    tipi = conn.execute(
        "SELECT id, nome FROM tipi_entita ORDER BY nome").fetchall()
    tipo_arg = request.args.get("tipo", "")
    tipo_id = int(tipo_arg) if tipo_arg.isdigit() else None
    righe = arc.completezza_entita(conn, tipo_id)
    conn.close()
    return render_template("archivio_completezza.html", righe=righe,
                           tipi=tipi, tipo_sel=tipo_arg,
                           tipologie_attese=arc.TIPOLOGIE_ATTESE)


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

@bp.route("/export/archivio_logico.json")
def export_json():
    conn = db.get_connection()
    dati = arc.costruisci_export(conn)
    conn.close()
    testo = json.dumps(dati, ensure_ascii=False, indent=2)
    return Response(testo, mimetype="application/json",
                    headers={"Content-Disposition":
                             "attachment; filename=archivio_logico.json"})


@bp.route("/export/archivio_logico.xlsx")
def export_xlsx():
    import openpyxl
    from openpyxl.styles import Font
    conn = db.get_connection()
    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Documenti"
    ws.append(["Percorso", "Cartella progetto", "Entità", "Tipologia",
               "Origine tipologia", "Riservatezza", "Origine riservatezza",
               "Stato"])
    for c in ws[1]:
        c.font = bold
    for d in conn.execute(
            "SELECT d.*, e.chiave AS entita_chiave FROM documenti d "
            "LEFT JOIN entita e ON e.id = d.entita_id "
            "ORDER BY d.cartella_progetto, d.percorso_rel"):
        ws.append([sanifica_cella(v) for v in (
            d["percorso_rel"], d["cartella_progetto"],
            d["entita_chiave"] or "", d["tipologia"], d["tipologia_origine"],
            d["riservatezza"] or "", d["riservatezza_origine"] or "",
            d["stato"])])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Regole di tipologia")
    ws2.append(["Priorità", "Nome", "Tipologia assegnata", "Attiva"])
    for c in ws2[1]:
        c.font = bold
    for r in arc.regole_tutte(conn):
        ws2.append([sanifica_cella(v) for v in (
            r["priorita"], r["nome"], r["tipologia"],
            "sì" if r["attiva"] else "no")])
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Completezza per entità")
    ws3.append(["Tipo entità", "Chiave", "Tipologie presenti",
               "Tipologie assenti", "% completezza"])
    for c in ws3[1]:
        c.font = bold
    for r in arc.completezza_entita(conn):
        ws3.append([sanifica_cella(v) for v in (
            r["entita"]["tipo_nome"], r["entita"]["chiave"],
            ", ".join(r["presenti"]), ", ".join(r["assenti"]),
            r["percentuale"])])
    ws3.freeze_panes = "A2"
    conn.close()

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(buf.read(),
                    mimetype="application/vnd.openxmlformats-officedocument"
                             ".spreadsheetml.sheet",
                    headers={"Content-Disposition":
                             "attachment; filename=archivio_logico.xlsx"})
