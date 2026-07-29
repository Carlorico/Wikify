# -*- coding: utf-8 -*-
"""
Modulo Catalogo: viste per la definizione delle entita' (2A) e per
l'arricchimento dati da file tabellari (2B). La logica di dominio vive in
core/catalogo.py.

Convenzione sugli endpoint: tutte le viste dell'arricchimento hanno un nome
che inizia per "arricchimento" (la sidebar distingue cosi' le due sotto-voci).
"""

import datetime
import io
import json
import os

from flask import (Blueprint, Response, flash, redirect, render_template,
                   request, session, url_for)

from core import catalogo as cat
from core import db
from core.report import sanifica_cella

bp = Blueprint("catalogo", __name__)

STATI_CHIARIMENTO = ["aperta", "chiarita", "falso_allarme"]
STATI_CHIARIMENTO_LABEL = {
    "aperta": "Aperta", "chiarita": "Chiarita", "falso_allarme": "Falso allarme"}

ANOMALIE_LABEL = {
    "chiave_senza_entita": "Chiave senza entità",
    "chiave_normalizzata": "Chiave normalizzata",
    "duplicato_esatto": "Duplicato esatto",
    "duplicato_normalizzazione": "Duplicato post-normalizzazione",
    "entita_non_citate": "Entità mai citate",
}


def _utente_sessione():
    return session.get("utente_nome") or ""


def _tipo_o_none(conn, tid):
    return conn.execute("SELECT * FROM tipi_entita WHERE id = ?",
                        (tid,)).fetchone()


# ---------------------------------------------------------------------------
# 2A - Definizione entita'
# ---------------------------------------------------------------------------

@bp.route("/")
def definizione():
    """Elenco dei tipi di entita' con lo stato attuale dell'estrazione."""
    conn = db.get_connection()
    tipi = []
    for t in conn.execute("SELECT * FROM tipi_entita ORDER BY nome"):
        criterio = json.loads(t["criterio_json"] or "{}")
        righe = cat.estrazioni(conn, criterio)
        conteggi = cat.riepilogo_estrazioni(righe)
        n_attive = conn.execute(
            "SELECT COUNT(*) FROM entita WHERE tipo_id = ? AND stato = 'attiva'",
            (t["id"],)).fetchone()[0]
        n_orfane = conn.execute(
            "SELECT COUNT(*) FROM entita WHERE tipo_id = ? "
            "AND stato = 'senza_riscontro'", (t["id"],)).fetchone()[0]
        tipi.append({"riga": t, "criterio": criterio, "conteggi": conteggi,
                     "n_attive": n_attive, "n_orfane": n_orfane,
                     "modo": cat.MODI.get(criterio.get("modo", ""), "")})
    n_cartelle = len(cat.cartelle_inventario(conn))
    conn.close()
    return render_template("catalogo_definizione.html", tipi=tipi,
                           n_cartelle=n_cartelle)


def _criterio_da_form(form, escluse=None):
    """Criterio (validato) dai campi del form. Restituisce (criterio, errore)."""
    modo = form.get("modo") or "intero"
    parametri = {}
    if modo == "prefisso":
        parametri = {k: (form.get(k) or "") for k in
                     ("prefisso", "blocco1_tipo", "blocco1_n",
                      "separatore", "blocco2_tipo", "blocco2_n")}
    elif modo == "regex":
        parametri = {"regex": form.get("regex") or ""}
    return cat.criterio_da_parametri(modo, parametri, escluse)


def _form_tipo(conn, tipo=None):
    """GET/POST comuni di creazione e modifica del tipo di entita'.

    Flusso: il pulsante "Prova sui nomi reali" (azione=prova) calcola la
    tabella cartella -> chiave; solo dopo la prova compare "Salva". Il
    salvataggio riesegue comunque la validazione e il riallineamento."""
    from modules.dizionario import CLASSI_LABEL

    escluse = []
    if tipo is not None:
        escluse = (json.loads(tipo["criterio_json"] or "{}")
                   .get("escluse") or [])
    prova, conteggi, errore = None, None, None
    if request.method == "POST":
        try:
            escluse = json.loads(request.form.get("escluse_json") or "[]")
        except ValueError:
            escluse = []
        azione = request.form.get("azione") or "prova"
        nome = (request.form.get("nome") or "").strip()
        criterio, errore = _criterio_da_form(request.form, escluse)
        if errore is None and not nome:
            errore = "Indicare il nome del tipo di entità (es. Progetto)."
        if errore is None:
            doppione = conn.execute(
                "SELECT id FROM tipi_entita WHERE nome = ?", (nome,)).fetchone()
            if doppione and (tipo is None or doppione["id"] != tipo["id"]):
                errore = "Esiste già un tipo di entità con questo nome."
        if errore is None:
            prova = cat.estrazioni(conn, criterio)
            conteggi = cat.riepilogo_estrazioni(prova)
            if azione == "salva":
                adesso = datetime.datetime.now().strftime("%d/%m/%Y %H:%M")
                if tipo is None:
                    cur = conn.execute(
                        "INSERT INTO tipi_entita (nome, criterio_json, "
                        "data_creazione, stato) VALUES (?, ?, ?, 'attivo')",
                        (nome, json.dumps(criterio), adesso))
                    tid = cur.lastrowid
                else:
                    conn.execute(
                        "UPDATE tipi_entita SET nome = ?, criterio_json = ? "
                        "WHERE id = ?", (nome, json.dumps(criterio), tipo["id"]))
                    tid = tipo["id"]
                conn.commit()
                riga = _tipo_o_none(conn, tid)
                report = cat.riallinea(conn, riga)
                conn.close()
                flash("Tipo di entità salvato: riallineamento eseguito.", "ok")
                return render_template("catalogo_riallineamento.html",
                                       tipo=riga, report=report)
        if errore:
            flash(errore, "errore")
    parametri, modo_default = {}, "intero"
    if tipo is not None and request.method == "GET":
        criterio_salvato = json.loads(tipo["criterio_json"] or "{}")
        parametri = criterio_salvato.get("parametri") or {}
        modo_default = criterio_salvato.get("modo") or "intero"
    conn.close()
    return render_template(
        "catalogo_tipo_form.html", tipo=tipo, prova=prova, conteggi=conteggi,
        classi=CLASSI_LABEL, modi=cat.MODI, escluse=escluse,
        parametri=parametri, modo_default=modo_default,
        form=request.form if request.method == "POST" else None)


@bp.route("/tipi/nuovo", methods=["GET", "POST"])
def tipo_nuovo():
    return _form_tipo(db.get_connection(), None)


@bp.route("/tipi/<int:tid>/modifica", methods=["GET", "POST"])
def tipo_modifica(tid):
    conn = db.get_connection()
    tipo = _tipo_o_none(conn, tid)
    if tipo is None:
        conn.close()
        flash("Tipo di entità non trovato.", "errore")
        return redirect(url_for("catalogo.definizione"))
    return _form_tipo(conn, tipo)


@bp.route("/tipi/<int:tid>/riallinea", methods=["POST"])
def tipo_riallinea(tid):
    """Riallineamento esplicito con report: crea le nuove, marca le orfane,
    non cancella mai."""
    conn = db.get_connection()
    tipo = _tipo_o_none(conn, tid)
    if tipo is None:
        conn.close()
        flash("Tipo di entità non trovato.", "errore")
        return redirect(url_for("catalogo.definizione"))
    report = cat.riallinea(conn, tipo)
    conn.close()
    return render_template("catalogo_riallineamento.html", tipo=tipo,
                           report=report)


@bp.route("/tipi/<int:tid>/escludi", methods=["POST"])
def tipo_escludi(tid):
    """Esclude una cartella dalla generazione (risoluzione manuale di un
    conflitto di chiave) e riesegue il riallineamento."""
    conn = db.get_connection()
    tipo = _tipo_o_none(conn, tid)
    cartella = (request.form.get("cartella") or "").strip()
    if tipo is None or not cartella:
        conn.close()
        flash("Richiesta di esclusione non valida.", "errore")
        return redirect(url_for("catalogo.definizione"))
    criterio = json.loads(tipo["criterio_json"] or "{}")
    escluse = set(criterio.get("escluse") or [])
    escluse.add(cartella)
    criterio["escluse"] = sorted(escluse)
    conn.execute("UPDATE tipi_entita SET criterio_json = ? WHERE id = ?",
                 (json.dumps(criterio), tid))
    conn.commit()
    tipo = _tipo_o_none(conn, tid)
    report = cat.riallinea(conn, tipo)
    conn.close()
    flash("Cartella esclusa dalla generazione: %s." % cartella, "ok")
    return render_template("catalogo_riallineamento.html", tipo=tipo,
                           report=report)


@bp.route("/entita")
def entita_lista():
    """Elenco entita' filtrabile per tipo, chiave e stato."""
    conn = db.get_connection()
    f_tipo = request.args.get("tipo", "")
    f_chiave = (request.args.get("q") or "").strip()
    f_stato = request.args.get("stato", "")
    clausole, valori = [], []
    if f_tipo.isdigit():
        clausole.append("e.tipo_id = ?")
        valori.append(int(f_tipo))
    if f_chiave:
        clausole.append("(e.chiave LIKE ? OR e.cartella_origine LIKE ?)")
        valori += ["%" + f_chiave + "%"] * 2
    if f_stato in ("attiva", "senza_riscontro"):
        clausole.append("e.stato = ?")
        valori.append(f_stato)
    where = (" WHERE " + " AND ".join(clausole)) if clausole else ""
    righe = conn.execute(
        "SELECT e.*, t.nome AS tipo_nome, "
        "(SELECT COUNT(*) FROM attributi_entita a "
        " JOIN importazioni i ON i.id = a.importazione_id "
        " WHERE a.entita_id = e.id AND i.stato = 'confermata') AS n_attributi "
        "FROM entita e JOIN tipi_entita t ON t.id = e.tipo_id"
        + where + " ORDER BY t.nome, e.chiave", valori).fetchall()
    tipi = conn.execute("SELECT id, nome FROM tipi_entita ORDER BY nome").fetchall()
    conn.close()
    return render_template("catalogo_entita.html", righe=righe, tipi=tipi,
                           f_tipo=f_tipo, f_chiave=f_chiave, f_stato=f_stato)


@bp.route("/entita/<int:eid>")
def entita_scheda(eid):
    """Scheda entita': chiave, cartella, stato, attributi con provenienza."""
    conn = db.get_connection()
    entita = conn.execute(
        "SELECT e.*, t.nome AS tipo_nome FROM entita e "
        "JOIN tipi_entita t ON t.id = e.tipo_id WHERE e.id = ?",
        (eid,)).fetchone()
    if entita is None:
        conn.close()
        flash("Entità non trovata.", "errore")
        return redirect(url_for("catalogo.entita_lista"))
    attributi = cat.attributi_di_entita(conn, eid)
    conn.close()
    return render_template("catalogo_entita_scheda.html", entita=entita,
                           attributi=attributi)


@bp.route("/export/catalogo.xlsx")
def export_catalogo():
    """Export XLSX del catalogo: entita' e attributi con provenienza."""
    import openpyxl
    from openpyxl.styles import Font
    conn = db.get_connection()
    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Entità"
    ws.append(["Tipo", "Chiave", "Cartella di origine", "Stato",
               "Data creazione", "Ultimo riallineamento", "N. attributi"])
    for c in ws[1]:
        c.font = bold
    entita = conn.execute(
        "SELECT e.*, t.nome AS tipo_nome, "
        "(SELECT COUNT(*) FROM attributi_entita a "
        " JOIN importazioni i ON i.id = a.importazione_id "
        " WHERE a.entita_id = e.id AND i.stato = 'confermata') AS n_attributi "
        "FROM entita e JOIN tipi_entita t ON t.id = e.tipo_id "
        "ORDER BY t.nome, e.chiave").fetchall()
    for e in entita:
        ws.append([sanifica_cella(v) for v in (
            e["tipo_nome"], e["chiave"], e["cartella_origine"],
            "attiva" if e["stato"] == "attiva" else "senza riscontro",
            e["data_creazione"], e["data_riallineamento"], e["n_attributi"])])
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Attributi")
    ws.append(["Tipo", "Chiave entità", "Attributo", "Valore",
               "Provenienza (import)"])
    for c in ws[1]:
        c.font = bold
    for e in entita:
        for a in cat.attributi_di_entita(conn, e["id"]):
            provenienza = "; ".join(
                "%s (%s, %s)" % (p["file"], p["data"], p["utente"])
                for p in a["provenienze"])
            ws.append([sanifica_cella(v) for v in (
                e["tipo_nome"], e["chiave"], a["attributo"], a["valore"],
                provenienza)])
    ws.freeze_panes = "A2"
    conn.close()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(buf.read(),
                    mimetype="application/vnd.openxmlformats-officedocument"
                             ".spreadsheetml.sheet",
                    headers={"Content-Disposition":
                             "attachment; filename=catalogo_entita.xlsx"})


# ---------------------------------------------------------------------------
# 2B - Arricchimento dati
# ---------------------------------------------------------------------------

@bp.route("/arricchimento")
def arricchimento():
    """Upload di un nuovo file ed elenco delle importazioni."""
    conn = db.get_connection()
    tipi = conn.execute("SELECT id, nome FROM tipi_entita ORDER BY nome").fetchall()
    importazioni = conn.execute(
        "SELECT i.*, t.nome AS tipo_nome, "
        "(SELECT COUNT(*) FROM anomalie_import an "
        " WHERE an.importazione_id = i.id "
        " AND an.stato_chiarimento = 'aperta') AS n_anomalie_aperte, "
        "(SELECT COUNT(*) FROM anomalie_import an "
        " WHERE an.importazione_id = i.id) AS n_anomalie "
        "FROM importazioni i JOIN tipi_entita t ON t.id = i.tipo_entita_id "
        "ORDER BY i.id DESC").fetchall()
    conn.close()
    return render_template("catalogo_arricchimento.html", tipi=tipi,
                           importazioni=importazioni)


def _percorso_temp(token):
    if not token.isalnum():
        return None
    return cat.cartella_temporanea() / (token + ".json")


def _carica_temp(token):
    percorso = _percorso_temp(token)
    if percorso is None or not percorso.is_file():
        return None
    try:
        return json.loads(percorso.read_text(encoding="utf-8"))
    except ValueError:
        return None


@bp.route("/arricchimento/carica", methods=["POST"])
def arricchimento_carica():
    """Passo 1: upload del file, lettura e salvataggio di lavoro."""
    f = request.files.get("file_dati")
    tipo_id = request.form.get("tipo_entita") or ""
    conn = db.get_connection()
    tipo = _tipo_o_none(conn, tipo_id) if tipo_id.isdigit() else None
    conn.close()
    if tipo is None:
        flash("Scegliere il tipo di entità a cui agganciare i dati.", "errore")
        return redirect(url_for("catalogo.arricchimento"))
    if f is None or not f.filename:
        flash("Selezionare un file CSV o XLSX da caricare.", "errore")
        return redirect(url_for("catalogo.arricchimento"))
    intestazioni, righe, errore = cat.leggi_tabella(f.filename, f.read())
    if errore:
        flash(errore, "errore")
        return redirect(url_for("catalogo.arricchimento"))
    token = os.urandom(8).hex()
    _percorso_temp(token).write_text(json.dumps({
        "nome_file": f.filename, "tipo_entita_id": tipo["id"],
        "intestazioni": intestazioni, "righe": righe}), encoding="utf-8")
    return redirect(url_for("catalogo.arricchimento_mappatura", token=token))


@bp.route("/arricchimento/<token>/mappatura")
def arricchimento_mappatura(token):
    """Passo 2: anteprima delle prime 10 righe e mappatura delle colonne."""
    dati = _carica_temp(token)
    if dati is None:
        flash("Sessione di caricamento non trovata: ripetere l'upload.", "errore")
        return redirect(url_for("catalogo.arricchimento"))
    conn = db.get_connection()
    tipo = _tipo_o_none(conn, dati["tipo_entita_id"])
    conn.close()
    return render_template("catalogo_mappatura.html", token=token, dati=dati,
                           tipo=tipo, anteprima=dati["righe"][:10])


def _mappatura_da_form(form, n_colonne):
    """Mappatura dal form del passo 2. Restituisce (mappatura, errore)."""
    chiave_col = form.get("chiave_col") or ""
    if not chiave_col.isdigit() or int(chiave_col) >= n_colonne:
        return None, "Indicare quale colonna contiene la chiave dell'entità."
    chiave_col = int(chiave_col)
    attributi, nomi_visti = {}, set()
    for i in range(n_colonne):
        if i == chiave_col or not form.get("includi_%d" % i):
            continue
        nome = (form.get("nome_%d" % i) or "").strip()
        if not nome:
            return None, "Indicare il nome dell'attributo per ogni colonna inclusa."
        if nome.lower() in nomi_visti:
            return None, "Due colonne incluse hanno lo stesso nome di attributo: %s." % nome
        nomi_visti.add(nome.lower())
        attributi[i] = nome
    if not attributi:
        return None, "Includere almeno una colonna come attributo."
    return {"chiave_col": chiave_col, "attributi": attributi}, None


@bp.route("/arricchimento/<token>/anteprima", methods=["POST"])
def arricchimento_anteprima(token):
    """Passo 3: anteprima di qualita' prima della conferma."""
    dati = _carica_temp(token)
    if dati is None:
        flash("Sessione di caricamento non trovata: ripetere l'upload.", "errore")
        return redirect(url_for("catalogo.arricchimento"))
    mappatura, errore = _mappatura_da_form(request.form,
                                           len(dati["intestazioni"]))
    if errore:
        flash(errore, "errore")
        return redirect(url_for("catalogo.arricchimento_mappatura",
                                token=token))
    conn = db.get_connection()
    tipo = _tipo_o_none(conn, dati["tipo_entita_id"])
    analisi = cat.analizza_import(conn, tipo["id"], dati["intestazioni"],
                                  dati["righe"], mappatura)
    conn.close()
    return render_template(
        "catalogo_anteprima.html", token=token, dati=dati, tipo=tipo,
        analisi=analisi, mappatura=mappatura,
        mappatura_json=json.dumps(mappatura),
        intestazioni_mappate=[
            (dati["intestazioni"][int(i)], nome)
            for i, nome in sorted(mappatura["attributi"].items())])


@bp.route("/arricchimento/<token>/conferma", methods=["POST"])
def arricchimento_conferma(token):
    """Passo 4: conferma e salvataggio con provenienza."""
    dati = _carica_temp(token)
    if dati is None:
        flash("Sessione di caricamento non trovata: ripetere l'upload.", "errore")
        return redirect(url_for("catalogo.arricchimento"))
    try:
        mappatura = json.loads(request.form.get("mappatura_json") or "")
        mappatura = {"chiave_col": int(mappatura["chiave_col"]),
                     "attributi": {int(k): str(v) for k, v in
                                   mappatura["attributi"].items()}}
    except (ValueError, KeyError, TypeError):
        flash("Mappatura non valida: ripetere il passaggio precedente.", "errore")
        return redirect(url_for("catalogo.arricchimento_mappatura",
                                token=token))
    conn = db.get_connection()
    tipo = _tipo_o_none(conn, dati["tipo_entita_id"])
    analisi = cat.analizza_import(conn, tipo["id"], dati["intestazioni"],
                                  dati["righe"], mappatura)
    imp_id = cat.esegui_import(conn, tipo["id"], dati["nome_file"],
                               mappatura, analisi, _utente_sessione())
    conn.close()
    percorso = _percorso_temp(token)
    if percorso is not None and percorso.is_file():
        percorso.unlink()
    flash("Importazione completata: %d valori salvati, %d righe scartate."
          % (analisi["n_valori"], len(analisi["scartate"])), "ok")
    return redirect(url_for("catalogo.arricchimento_importazione",
                            imp_id=imp_id))


@bp.route("/arricchimento/importazioni/<int:imp_id>")
def arricchimento_importazione(imp_id):
    """Dettaglio importazione: parametri, valori e anomalie gestibili."""
    conn = db.get_connection()
    imp = conn.execute(
        "SELECT i.*, t.nome AS tipo_nome FROM importazioni i "
        "JOIN tipi_entita t ON t.id = i.tipo_entita_id WHERE i.id = ?",
        (imp_id,)).fetchone()
    if imp is None:
        conn.close()
        flash("Importazione non trovata.", "errore")
        return redirect(url_for("catalogo.arricchimento"))
    anomalie = conn.execute(
        "SELECT * FROM anomalie_import WHERE importazione_id = ? "
        "ORDER BY id", (imp_id,)).fetchall()
    valori = conn.execute(
        "SELECT a.nome_attributo, a.valore, e.chiave FROM attributi_entita a "
        "JOIN entita e ON e.id = a.entita_id "
        "WHERE a.importazione_id = ? "
        "ORDER BY e.chiave, a.nome_attributo, a.valore", (imp_id,)).fetchall()
    conn.close()
    return render_template("catalogo_importazione.html", imp=imp,
                           anomalie=anomalie, valori=valori,
                           stati=STATI_CHIARIMENTO,
                           stati_label=STATI_CHIARIMENTO_LABEL,
                           anomalie_label=ANOMALIE_LABEL)


@bp.route("/arricchimento/importazioni/<int:imp_id>/rimuovi", methods=["POST"])
def arricchimento_rimuovi(imp_id):
    """Rimozione in blocco dei valori dell'import (rollback per provenienza)."""
    conn = db.get_connection()
    imp = conn.execute("SELECT * FROM importazioni WHERE id = ?",
                       (imp_id,)).fetchone()
    if imp is None:
        conn.close()
        flash("Importazione non trovata.", "errore")
        return redirect(url_for("catalogo.arricchimento"))
    n = cat.rimuovi_import(conn, imp_id)
    conn.close()
    flash("Importazione rimossa: %d valori eliminati. I valori identici "
          "portati da altri import restano disponibili." % n, "ok")
    return redirect(url_for("catalogo.arricchimento"))


@bp.route("/arricchimento/anomalie/<int:aid>/stato", methods=["POST"])
def arricchimento_anomalia_stato(aid):
    """Aggiorna lo stato di chiarimento di un'anomalia."""
    stato = request.form.get("stato_chiarimento") or ""
    nota = (request.form.get("nota_chiarimento") or "").strip()
    conn = db.get_connection()
    riga = conn.execute("SELECT * FROM anomalie_import WHERE id = ?",
                        (aid,)).fetchone()
    if riga is None or stato not in STATI_CHIARIMENTO:
        conn.close()
        flash("Anomalia non trovata o stato non valido.", "errore")
        return redirect(url_for("catalogo.arricchimento"))
    conn.execute("UPDATE anomalie_import SET stato_chiarimento = ?, "
                 "nota_chiarimento = ? WHERE id = ?", (stato, nota, aid))
    conn.commit()
    conn.close()
    flash("Anomalia aggiornata.", "ok")
    return redirect(url_for("catalogo.arricchimento_importazione",
                            imp_id=riga["importazione_id"]))


@bp.route("/arricchimento/export/anomalie.xlsx")
def arricchimento_export_anomalie():
    """Export XLSX del report anomalie di tutti gli import."""
    import openpyxl
    from openpyxl.styles import Font
    conn = db.get_connection()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Anomalie import"
    ws.append(["Import", "File", "Data", "Tipo entità", "Tipo anomalia",
               "Riferimento", "Dettaglio", "Stato chiarimento", "Nota"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in conn.execute(
            "SELECT an.*, i.nome_file, i.data AS data_import, "
            "t.nome AS tipo_nome "
            "FROM anomalie_import an "
            "JOIN importazioni i ON i.id = an.importazione_id "
            "JOIN tipi_entita t ON t.id = i.tipo_entita_id "
            "ORDER BY an.importazione_id, an.id"):
        ws.append([sanifica_cella(v) for v in (
            r["importazione_id"], r["nome_file"], r["data_import"],
            r["tipo_nome"], ANOMALIE_LABEL.get(r["tipo"], r["tipo"]),
            r["riferimento"], r["dettaglio"],
            STATI_CHIARIMENTO_LABEL.get(r["stato_chiarimento"],
                                        r["stato_chiarimento"]),
            r["nota_chiarimento"])])
    ws.freeze_panes = "A2"
    conn.close()
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(buf.read(),
                    mimetype="application/vnd.openxmlformats-officedocument"
                             ".spreadsheetml.sheet",
                    headers={"Content-Disposition":
                             "attachment; filename=anomalie_import.xlsx"})
