# -*- coding: utf-8 -*-
"""
Livello 0: articoli come entita' e relazione famiglia/articolo.

Primo blocco della base deterministica (progettazione in
!!IINTERNO!!!/modello_dati_livello_0.md). Fa tre cose, tutte senza
interpretazione del contenuto dei documenti:

  1. deriva le famiglie dall'albero dell'inventario, con una regola che non
     dipende da un livello fisso di profondita';
  2. legge il listino e crea un'entita' per ogni codice articolo, ricavando
     la relazione famiglia/articolo dalle righe di sezione, che sono l'unica
     rappresentazione esplicita di quel legame;
  3. misura la copertura, cioe' quante famiglie trovano riscontro a listino e
     quante no, in entrambe le direzioni.

Ogni abbinamento porta con se' il modo in cui e' stato ottenuto; cio' che non
si abbina non viene forzato, diventa un'anomalia da leggere.
"""

import datetime
import json
import os
import re

from core import db

TIPO_FAMIGLIA = "Famiglia"
TIPO_ARTICOLO = "Articolo"
RELAZIONE = "raccoglie"

# Nomi di cartella che denotano una tipologia documentale e non un prodotto.
# Sono il marcatore che consente di riconoscere una famiglia senza dipendere
# dal livello in cui si trova.
MARCATORI_TIPOLOGIA = ("DATA SHEETS", "OPERATION MANUALS", "CERTIFICATES",
                       "DRAWINGS", "SOFTWARE", "SCHEDE TECNICHE", "MANUALI")

CHIAVE_MARCATORI = "livello0_marcatori_tipologia"


def _adesso():
    return datetime.datetime.now().strftime("%d/%m/%Y %H:%M")


def _testo(valore):
    """Valore di cella come testo pulito: openpyxl restituisce numeri per i
    codici tutti numerici, e un codice non e' un numero."""
    if valore is None:
        return ""
    if isinstance(valore, float) and valore.is_integer():
        return str(int(valore))
    return str(valore).strip()


def _norm(testo):
    return re.sub(r"\s+", " ", (testo or "").strip().upper())


def _senza_spazi(testo):
    return _norm(testo).replace(" ", "")


def marcatori(conn):
    """Marcatori di tipologia, sovrascrivibili da impostazione."""
    salvati = db.get_impostazione(conn, CHIAVE_MARCATORI, "")
    if salvati.strip():
        return tuple(_norm(x) for x in salvati.split("|") if x.strip())
    return tuple(_norm(x) for x in MARCATORI_TIPOLOGIA)


# ===========================================================================
# 1. Le famiglie, derivate dall'albero
# ===========================================================================

def famiglie_da_inventario(conn, elenco_marcatori=None):
    """Riconosce le famiglie fra le cartelle dell'inventario.

    Regola dichiarata, indipendente dal livello di profondita':

        una cartella e' una famiglia se sta almeno al secondo livello sotto
        la radice, il suo nome non e' un marcatore di tipologia, e o contiene
        una sottocartella marcatore, oppure contiene direttamente documenti.

    Il primo livello e' la categoria commerciale e non e' mai una famiglia;
    le cartelle di servizio e quelle di catalogo, che stanno al primo livello
    e contengono documenti non riferibili a un prodotto, restano fuori per
    costruzione.

    Restituisce una lista ordinata di dizionari {chiave, percorso, livello,
    con_marcatore}."""
    marca = set(elenco_marcatori or marcatori(conn))
    con_file = set()        # cartelle che contengono direttamente documenti
    figli_marcatori = {}    # cartella -> insieme dei figli marcatore
    for (percorso_rel,) in conn.execute(
            "SELECT percorso_rel FROM inventario_file WHERE stato = 'presente'"):
        parti = percorso_rel.replace("\\", "/").split("/")
        if len(parti) < 2:
            continue
        cartelle = parti[:-1]
        con_file.add("/".join(cartelle))
        for i in range(1, len(cartelle)):
            genitore = "/".join(cartelle[:i])
            figlio = cartelle[i]
            if _norm(figlio) in marca:
                figli_marcatori.setdefault(genitore, set()).add(figlio)

    candidate = set(con_file) | set(figli_marcatori)
    fuori = []
    for percorso in sorted(candidate):
        parti = percorso.split("/")
        nome = parti[-1]
        if len(parti) < 2:
            continue                      # primo livello: categoria
        if _norm(nome) in marca:
            continue                      # la cartella stessa e' una tipologia
        if percorso not in figli_marcatori and percorso not in con_file:
            continue
        fuori.append({"chiave": nome, "percorso": percorso,
                      "livello": len(parti),
                      "con_marcatore": percorso in figli_marcatori})
    return fuori


def assicura_tipo(conn, nome, criterio=None):
    riga = conn.execute("SELECT id FROM tipi_entita WHERE nome = ?",
                        (nome,)).fetchone()
    if riga:
        return riga[0]
    cur = conn.execute(
        "INSERT INTO tipi_entita (nome, criterio_json, data_creazione, stato) "
        "VALUES (?, ?, ?, 'attivo')",
        (nome, json.dumps(criterio or {}, ensure_ascii=False), _adesso()))
    return cur.lastrowid


def sincronizza_famiglie(conn, elenco_marcatori=None):
    """Crea o riallinea le entita' di tipo Famiglia. Non cancella mai: le
    famiglie non piu' riscontrate vengono marcate 'senza riscontro', come
    gia' avviene nel catalogo."""
    trovate = famiglie_da_inventario(conn, elenco_marcatori)
    tipo_id = assicura_tipo(conn, TIPO_FAMIGLIA, {
        "sorgente": "albero_inventario",
        "regola": "cartella dal secondo livello, con sottocartella marcatore "
                  "oppure con documenti diretti",
        "marcatori": list(elenco_marcatori or marcatori(conn))})

    per_chiave = {}
    conflitti = []
    for f in trovate:
        if f["chiave"] in per_chiave:
            conflitti.append({"chiave": f["chiave"],
                              "percorsi": [per_chiave[f["chiave"]]["percorso"],
                                           f["percorso"]]})
            continue
        per_chiave[f["chiave"]] = f

    esistenti = {r["chiave"]: r["id"] for r in conn.execute(
        "SELECT id, chiave FROM entita WHERE tipo_id = ?", (tipo_id,))}
    nuove = 0
    for chiave, f in sorted(per_chiave.items()):
        if chiave in esistenti:
            conn.execute("UPDATE entita SET cartella_origine = ?, "
                         "stato = 'attiva', data_riallineamento = ? "
                         "WHERE id = ?",
                         (f["percorso"], _adesso(), esistenti[chiave]))
        else:
            conn.execute(
                "INSERT INTO entita (tipo_id, chiave, cartella_origine, "
                "origine, stato, data_creazione, data_riallineamento) "
                "VALUES (?, ?, ?, 'cartella', 'attiva', ?, ?)",
                (tipo_id, chiave, f["percorso"], _adesso(), _adesso()))
            nuove += 1
    scomparse = sorted(set(esistenti) - set(per_chiave))
    for chiave in scomparse:
        conn.execute("UPDATE entita SET stato = 'senza riscontro', "
                     "data_riallineamento = ? WHERE id = ?",
                     (_adesso(), esistenti[chiave]))
    conn.commit()
    return {"tipo_id": tipo_id, "trovate": len(per_chiave), "nuove": nuove,
            "scomparse": scomparse, "conflitti": conflitti,
            "famiglie": sorted(per_chiave.values(),
                               key=lambda x: x["percorso"])}


# ===========================================================================
# 2. Il listino, e la relazione che nasce dalle righe di sezione
# ===========================================================================

def leggi_listino(percorso):
    """Legge il listino multi foglio in una lista di righe tipizzate.

    Ogni foglio puo' avere l'intestazione su una riga diversa: viene cercata
    e non presunta. Una riga priva di codice e con descrizione apre un
    contesto di famiglia; le righe con codice che seguono appartengono a
    quel contesto."""
    import openpyxl

    wb = openpyxl.load_workbook(percorso, read_only=True, data_only=True)
    righe, fogli = [], {}
    for nome_foglio in wb.sheetnames:
        ws = wb[nome_foglio]
        intestazione = None
        col_codice = col_desc = None
        for n, riga in enumerate(ws.iter_rows(values_only=True), start=1):
            valori = [_testo(v) for v in riga]
            if intestazione is None:
                minuscoli = [v.lower() for v in valori]
                if "code" in minuscoli:
                    intestazione = n
                    col_codice = minuscoli.index("code")
                    col_desc = (minuscoli.index("description")
                                if "description" in minuscoli
                                else col_codice + 1)
                if n > 12:
                    break           # nessuna intestazione riconoscibile
                continue
            codice = valori[col_codice] if col_codice < len(valori) else ""
            desc = valori[col_desc] if col_desc < len(valori) else ""
            if not codice and not desc:
                continue
            righe.append({"foglio": nome_foglio, "riga": n,
                          "tipo": "articolo" if codice else "sezione",
                          "codice": codice, "descrizione": desc})
        fogli[nome_foglio] = {"intestazione_riga": intestazione,
                              "colonna_codice": col_codice,
                              "colonna_descrizione": col_desc}
    wb.close()
    return righe, fogli


def abbina_famiglia(testo_sezione, chiavi):
    """Abbina il testo di una riga di sezione a una chiave di famiglia.

    Restituisce (chiave, esito) con esito in 'esatto', 'normalizzato',
    'ambigua', 'nessuna'. Due passaggi dichiarati: confronto sul testo
    normalizzato, poi confronto ignorando gli spazi, che e' il caso della
    sigla scritta in due varianti. Se il resto della sezione riprende con una
    cifra l'abbinamento e' rifiutato, perche' significherebbe aver spezzato
    un numero a meta'."""
    s = _norm(testo_sezione)
    esatti = [k for k in chiavi
              if s == _norm(k) or s.startswith(_norm(k) + " ")]
    if esatti:
        lunghezza = max(len(_norm(k)) for k in esatti)
        vincitori = [k for k in esatti if len(_norm(k)) == lunghezza]
        if len(vincitori) > 1:
            return None, "ambigua"
        return vincitori[0], "esatto"

    sz = _senza_spazi(testo_sezione)
    candidati = []
    for k in chiavi:
        kz = _senza_spazi(k)
        if not kz or not sz.startswith(kz):
            continue
        resto = sz[len(kz):]
        if resto[:1].isdigit():
            continue
        candidati.append(k)
    if not candidati:
        return None, "nessuna"
    lunghezza = max(len(_senza_spazi(k)) for k in candidati)
    vincitori = [k for k in candidati if len(_senza_spazi(k)) == lunghezza]
    if len(vincitori) > 1:
        return None, "ambigua"
    return vincitori[0], "normalizzato"


def importa_listino(conn, percorso, utente="", sostituisci=True):
    """Crea le entita' di tipo Articolo e la relazione con le famiglie.

    Rieseguibile: un import precedente dello stesso file viene rimosso con
    tutti i valori e le relazioni che aveva portato, cosi' la ricostruzione
    e' idempotente e non lascia sedimenti."""
    nome_file = os.path.basename(percorso)
    tipo_articolo = assicura_tipo(conn, TIPO_ARTICOLO,
                                  {"sorgente": "listino", "chiave": "codice"})
    tipo_famiglia = assicura_tipo(conn, TIPO_FAMIGLIA)

    if sostituisci:
        vecchi = [r[0] for r in conn.execute(
            "SELECT id FROM importazioni WHERE nome_file = ? AND "
            "tipo_entita_id = ?", (nome_file, tipo_articolo))]
        for imp in vecchi:
            conn.execute("DELETE FROM attributi_entita WHERE importazione_id = ?",
                         (imp,))
            conn.execute("DELETE FROM relazioni_entita WHERE importazione_id = ?",
                         (imp,))
            conn.execute("DELETE FROM anomalie_import WHERE importazione_id = ?",
                         (imp,))
            conn.execute("DELETE FROM importazioni WHERE id = ?", (imp,))

    righe, fogli = leggi_listino(percorso)
    cur = conn.execute(
        "INSERT INTO importazioni (data, nome_file, tipo_entita_id, "
        "mappatura_json, n_righe, utente, stato) "
        "VALUES (?, ?, ?, ?, ?, ?, 'confermata')",
        (_adesso(), nome_file, tipo_articolo,
         json.dumps(fogli, ensure_ascii=False), len(righe), utente))
    imp_id = cur.lastrowid

    famiglie = {r["chiave"]: r["id"] for r in conn.execute(
        "SELECT id, chiave FROM entita WHERE tipo_id = ? AND stato = 'attiva'",
        (tipo_famiglia,))}
    articoli_esistenti = {r["chiave"]: r["id"] for r in conn.execute(
        "SELECT id, chiave FROM entita WHERE tipo_id = ?", (tipo_articolo,))}

    contesto = None          # (chiave_famiglia, esito) della sezione corrente
    visti = {}
    conteggi = {"sezioni": 0, "sezioni_abbinate": 0, "righe_articolo": 0,
                "articoli_agganciati": 0, "articoli_orfani": 0,
                "nuovi_articoli": 0, "relazioni": 0}
    anomalie = []

    def anomalia(tipo, riferimento, dettaglio):
        anomalie.append((tipo, riferimento, dettaglio))
        conn.execute("INSERT INTO anomalie_import (importazione_id, tipo, "
                     "riferimento, dettaglio) VALUES (?, ?, ?, ?)",
                     (imp_id, tipo, riferimento, dettaglio))

    for r in righe:
        if r["tipo"] == "sezione":
            conteggi["sezioni"] += 1
            chiave, esito = abbina_famiglia(r["descrizione"], famiglie.keys())
            contesto = (chiave, esito)
            if chiave:
                conteggi["sezioni_abbinate"] += 1
            else:
                anomalia("sezione_senza_famiglia",
                         "%s!%d" % (r["foglio"], r["riga"]),
                         "sezione '%s': %s" % (r["descrizione"], esito))
            continue

        conteggi["righe_articolo"] += 1
        codice = _norm(r["codice"])
        if codice in visti:
            if visti[codice] != r["descrizione"]:
                anomalia("codice_duplicato_divergente",
                         codice,
                         "descrizione diversa in %s!%d; conservata la prima"
                         % (r["foglio"], r["riga"]))
            continue
        visti[codice] = r["descrizione"]

        art_id = articoli_esistenti.get(codice)
        if art_id is None:
            art_id = conn.execute(
                "INSERT INTO entita (tipo_id, chiave, cartella_origine, "
                "origine, stato, data_creazione, data_riallineamento) "
                "VALUES (?, ?, '', 'import', 'attiva', ?, ?)",
                (tipo_articolo, codice, _adesso(), _adesso())).lastrowid
            articoli_esistenti[codice] = art_id
            conteggi["nuovi_articoli"] += 1
        else:
            conn.execute("UPDATE entita SET stato = 'attiva', "
                         "data_riallineamento = ? WHERE id = ?",
                         (_adesso(), art_id))
        if r["descrizione"]:
            conn.execute(
                "INSERT OR IGNORE INTO attributi_entita (entita_id, "
                "nome_attributo, valore, importazione_id, data, origine, "
                "visibilita) VALUES (?, 'descrizione_listino', ?, ?, ?, "
                "'import', 'pubblica')",
                (art_id, r["descrizione"], imp_id, _adesso()))

        if contesto and contesto[0]:
            conn.execute(
                "INSERT OR IGNORE INTO relazioni_entita (tipo_relazione, "
                "entita_da_id, entita_a_id, origine, importazione_id, data) "
                "VALUES (?, ?, ?, 'sezione_listino', ?, ?)",
                (RELAZIONE, famiglie[contesto[0]], art_id, imp_id, _adesso()))
            conteggi["articoli_agganciati"] += 1
            conteggi["relazioni"] += 1
        else:
            conteggi["articoli_orfani"] += 1
            anomalia("articolo_senza_famiglia", codice,
                     "nessuna sezione di famiglia abbinata sopra la riga "
                     "%s!%d" % (r["foglio"], r["riga"]))

    conn.execute("UPDATE importazioni SET n_valori = ?, n_scartate = ? "
                 "WHERE id = ?",
                 (conteggi["righe_articolo"], len(anomalie), imp_id))
    conn.commit()
    return {"importazione_id": imp_id, "fogli": fogli, "conteggi": conteggi,
            "anomalie": anomalie}


# ===========================================================================
# 3. La misura
# ===========================================================================

def misura_copertura(conn):
    """La copertura in entrambe le direzioni: famiglie senza articoli e
    articoli senza famiglia. E' il primo indicatore del livello 0, e va letto
    come una misura, non come un'impressione."""
    # La forma del risultato e' sempre la stessa, anche quando un tipo di
    # entita' non esiste ancora perche' il listino non e' stato importato:
    # una funzione che restituisce due forme diverse fa fallire chi la legge
    # proprio nel caso in cui l'archivio e' appena stato aperto.
    def _tipo(nome):
        riga = conn.execute("SELECT id FROM tipi_entita WHERE nome = ?",
                            (nome,)).fetchone()
        return riga[0] if riga else None

    def _entita(tipo_id, verso):
        if tipo_id is None:
            return []
        colonna = "entita_da_id" if verso == "da" else "entita_a_id"
        return [dict(r) for r in conn.execute(
            "SELECT e.id, e.chiave, e.cartella_origine, "
            "  (SELECT COUNT(*) FROM relazioni_entita r "
            "   WHERE r.tipo_relazione = ? AND r.%s = e.id) AS n_collegate "
            "FROM entita e WHERE e.tipo_id = ? AND e.stato = 'attiva' "
            "ORDER BY e.chiave" % colonna, (RELAZIONE, tipo_id))]

    famiglie = [dict(f, n_articoli=f["n_collegate"])
                for f in _entita(_tipo(TIPO_FAMIGLIA), "da")]
    articoli = [dict(a, n_famiglie=a["n_collegate"])
                for a in _entita(_tipo(TIPO_ARTICOLO), "a")]

    senza_articoli = [f for f in famiglie if not f["n_articoli"]]
    senza_famiglia = [a for a in articoli if not a["n_famiglie"]]
    return {
        "famiglie": len(famiglie),
        "famiglie_con_articoli": len(famiglie) - len(senza_articoli),
        "famiglie_senza_articoli": senza_articoli,
        "articoli": len(articoli),
        "articoli_con_famiglia": len(articoli) - len(senza_famiglia),
        "articoli_senza_famiglia": senza_famiglia,
        "copertura_famiglie": (
            0.0 if not famiglie
            else round(100.0 * (len(famiglie) - len(senza_articoli))
                       / len(famiglie), 1)),
        "copertura_articoli": (
            0.0 if not articoli
            else round(100.0 * (len(articoli) - len(senza_famiglia))
                       / len(articoli), 1)),
    }


def ricostruisci(conn, radice, listino=None, utente=""):
    """Catena completa del livello 0, nell'ordine imposto dalle dipendenze.

    Inventario, famiglie dall'albero, articoli e relazione dal listino,
    censimento dei documenti, derivazione di lingua, date e ruolo, aggancio
    alle entita', calcolo di vigenza. Ogni passo produce una misura: nessuno
    di essi richiede servizi esterni o giudizio."""
    from core import archivio, inventario

    if not conn.execute("SELECT COUNT(*) FROM inventario_file").fetchone()[0]:
        inventario.inventario_iniziale(conn, radice)
    else:
        inventario.esegui_check(conn, radice)
    db.set_impostazione(conn, "radice_archivio", radice)

    fam = sincronizza_famiglie(conn)
    imp = importa_listino(conn, listino, utente=utente) if listino else None
    archivio.ricostruisci(conn)          # censimento dei documenti
    regole = applica_regole_documenti(conn, radice)
    legami = collega_documenti_entita(conn, radice)
    vigenza = calcola_vigenza(conn)
    coerenza = analizza_coerenza_date(conn)
    return {"famiglie": fam, "listino": imp,
            "copertura": misura_copertura(conn), "documenti": regole,
            "legami": legami, "vigenza": vigenza, "coerenza": coerenza,
            "gruppi_incompleti": gruppi_traduzione_incompleti(conn)}


# ===========================================================================
# 4. Lingua, gruppo di traduzione, data, revisione, ruolo temporale
# ===========================================================================
#
# Tutto cio' che segue si ricava da regole dichiarate e riproducibili. La data
# di modifica del filesystem non viene mai usata come data del documento:
# sull'archivio reale e' l'impronta di una migrazione massiva e non l'istante
# in cui il documento e' stato emesso.

LINGUE_PREDEFINITE = ("ITA", "ENG", "SPA", "FRA", "DEU")
CHIAVE_LINGUE = "livello0_lingue"
CHIAVE_LINGUA_BASE = "livello0_lingua_predefinita"

RUOLO_DESCRITTIVO = "descrittivo"
RUOLO_AGGIORNAMENTO = "aggiornamento"
RUOLO_NESSUNO = "nessuno"

# Attenzione: \b non e' un confine accanto all'underscore, che nei nomi di
# file e' invece un separatore a tutti gli effetti. I confini sono espressi
# come assenza di cifra o di lettera, non come \b.
RX_DATA_ISO = re.compile(r"(?<!\d)(\d{4})-(\d{2})-(\d{2})(?!\d)")
RX_REVISIONE = re.compile(
    r"(?i)(?<![0-9A-Za-z])(rev|ed)\.?\s?(\d{1,3})(?![0-9A-Za-z])")

REGOLE_SEME = [
    ("Lingua dal nome del file", "nome_file", "lingua", 10,
     {"tipo": "lingua_da_nome"}),
    ("Revisione dal nome del file", "nome_file", "revisione", 20,
     {"tipo": "revisione_da_nome"}),
    ("Data dal nome del file", "nome_file", "data_documento", 30,
     {"tipo": "data_da_nome"}),
    ("Data dichiarata nel documento", "contenuto", "data_documento", 40,
     {"tipo": "data_da_contenuto",
      "etichette": ["Data di emissione", "Data di rilascio"]}),
    ("Nota tecnica dal nome del file", "nome_file", "ruolo_temporale", 50,
     {"tipo": "ruolo_da_nome", "prefissi": ["TN."],
      "valore": RUOLO_AGGIORNAMENTO}),
    ("Documento di catalogo o certificazione", "percorso", "ruolo_temporale",
     60, {"tipo": "ruolo_da_percorso",
          "cartelle": ["PRODUCT CATALOGUE", "_PRODUCT CERTIFICATES"],
          "valore": RUOLO_NESSUNO}),
    ("Prodotti dichiarati nel documento", "contenuto", "entita_collegate", 70,
     {"tipo": "famiglie_da_contenuto", "etichette": ["Prodotti interessati"]}),
]


def lingue(conn):
    salvate = db.get_impostazione(conn, CHIAVE_LINGUE, "")
    if salvate.strip():
        return tuple(_norm(x) for x in salvate.split("|") if x.strip())
    return LINGUE_PREDEFINITE


def seed_regole_estrazione(conn):
    """Popola il dizionario delle regole di estrazione al primo avvio.
    Non sovrascrive mai regole esistenti: e' un seme, non una riscrittura."""
    if conn.execute("SELECT COUNT(*) FROM regole_estrazione").fetchone()[0]:
        return 0
    for nome, ambito, attributo, priorita, criterio in REGOLE_SEME:
        conn.execute(
            "INSERT INTO regole_estrazione (nome, ambito, attributo, "
            "criterio_json, priorita, attiva, data_creazione) "
            "VALUES (?, ?, ?, ?, ?, 1, ?)",
            (nome, ambito, attributo,
             json.dumps(criterio, ensure_ascii=False), priorita, _adesso()))
    conn.commit()
    return len(REGOLE_SEME)


def regole_estrazione_attive(conn):
    return [dict(r, criterio=json.loads(r["criterio_json"] or "{}"))
            for r in conn.execute(
                "SELECT * FROM regole_estrazione WHERE attiva = 1 "
                "ORDER BY priorita, id")]


def _token(nome):
    """Nome del file scomposto nei suoi elementi, senza estensione."""
    base = os.path.splitext(os.path.basename(nome or ""))[0]
    return [t for t in re.split(r"[\s_\-.,;()]+", base) if t]


def _gruppo_traduzione(nome, marcatore):
    """Chiave del gruppo di traduzione: il nome privato del marcatore di
    lingua e normalizzato. Due documenti che condividono la chiave sono la
    stessa cosa in lingue diverse. Non si accorpa mai per somiglianza: se il
    corpo del nome cambia da una lingua all'altra, il gruppo resta separato e
    il residuo va segnalato."""
    senza_data = RX_DATA_ISO.sub(" ", os.path.splitext(nome or "")[0])
    parti = [t for t in _token(senza_data + ".x") if t != marcatore]
    return " ".join(parti).upper()


def _data_da_testo(testo):
    m = RX_DATA_ISO.search(testo or "")
    if not m:
        return ""
    try:
        datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return ""
    return m.group(0)


def _testo_documento(radice, percorso_rel, limite_righe=40):
    """Prime righe del testo del documento. Serve alle sole regole di ambito
    'contenuto', che cercano un'etichetta dichiarata: non e' lettura
    interpretativa, e' ricerca di un marcatore."""
    from core import estrazione
    percorso = os.path.join(radice, percorso_rel.replace("/", os.sep))
    if not os.path.isfile(percorso):
        return []
    try:
        parti = estrazione.estrai(percorso)
    except Exception:
        return []
    return [t for _, t in parti[:limite_righe]]


def _valore_etichetta(righe, etichette):
    for riga in righe:
        for etichetta in etichette:
            pos = riga.lower().find(etichetta.lower())
            if pos >= 0:
                return riga[pos + len(etichetta):].lstrip(" :\t")
    return ""


def applica_regole_documenti(conn, radice):
    """Deriva lingua, gruppo di traduzione, data, revisione e ruolo temporale
    per ogni documento censito. Rispetta la precedenza: cio' che ha origine
    'manuale' non viene mai riscritto."""
    seed_regole_estrazione(conn)
    regole = regole_estrazione_attive(conn)
    codici = set(lingue(conn))
    lingua_base = db.get_impostazione(conn, CHIAVE_LINGUA_BASE, "")
    serve_contenuto = any(r["ambito"] == "contenuto" and
                          r["criterio"].get("tipo") != "famiglie_da_contenuto"
                          for r in regole)

    conteggi = {"documenti": 0, "con_lingua": 0, "lingua_da_nome": 0,
                "con_data": 0, "data_da_nome": 0, "data_da_contenuto": 0,
                "senza_data": 0, "con_revisione": 0,
                "aggiornamento": 0, "nessuno": 0}
    adesso = _adesso()

    for doc in conn.execute("SELECT id, percorso_rel, lingua_origine, "
                            "data_origine, ruolo_origine FROM documenti "
                            "WHERE stato = 'attivo'").fetchall():
        conteggi["documenti"] += 1
        nome = os.path.basename(doc["percorso_rel"])
        campi = {"nome_file": nome, "percorso": doc["percorso_rel"]}
        righe = None

        lingua, lingua_origine = "", ""
        revisione = ""
        data, data_origine = "", ""
        ruolo, ruolo_origine = RUOLO_DESCRITTIVO, "predefinito"

        for regola in regole:
            tipo = regola["criterio"].get("tipo")
            if tipo == "lingua_da_nome":
                trovate = [t for t in _token(nome) if t in codici]
                if trovate:
                    lingua, lingua_origine = trovate[-1], "nome"
            elif tipo == "revisione_da_nome":
                m = RX_REVISIONE.search(nome)
                if m:
                    revisione = (m.group(1) + m.group(2)).lower()
            elif tipo == "data_da_nome" and not data:
                trovata = _data_da_testo(nome)
                if trovata:
                    data, data_origine = trovata, "nome"
            elif tipo == "data_da_contenuto" and not data:
                if righe is None and serve_contenuto:
                    righe = _testo_documento(radice, doc["percorso_rel"])
                valore = _valore_etichetta(righe or [],
                                           regola["criterio"].get("etichette") or [])
                trovata = _data_da_testo(valore)
                if trovata:
                    data, data_origine = trovata, "contenuto"
            elif tipo == "ruolo_da_nome":
                for prefisso in regola["criterio"].get("prefissi") or []:
                    if nome.upper().startswith(prefisso.upper()):
                        ruolo = regola["criterio"].get("valore",
                                                       RUOLO_AGGIORNAMENTO)
                        ruolo_origine = "regola"
            elif tipo == "ruolo_da_percorso":
                primo = campi["percorso"].split("/")[0]
                if _norm(primo) in {_norm(c) for c in
                                    (regola["criterio"].get("cartelle") or [])}:
                    ruolo = regola["criterio"].get("valore", RUOLO_NESSUNO)
                    ruolo_origine = "regola"

        if not lingua and lingua_base:
            lingua, lingua_origine = _norm(lingua_base), "predefinita"
        gruppo = _gruppo_traduzione(nome, lingua) if lingua else ""

        if lingua:
            conteggi["con_lingua"] += 1
            if lingua_origine == "nome":
                conteggi["lingua_da_nome"] += 1
        if data:
            conteggi["con_data"] += 1
            conteggi["data_da_%s" % data_origine] += 1
        else:
            conteggi["senza_data"] += 1
        if revisione:
            conteggi["con_revisione"] += 1
        if ruolo == RUOLO_AGGIORNAMENTO:
            conteggi["aggiornamento"] += 1
        elif ruolo == RUOLO_NESSUNO:
            conteggi["nessuno"] += 1

        conn.execute(
            "UPDATE documenti SET "
            "lingua = CASE WHEN lingua_origine = 'manuale' THEN lingua ELSE ? END, "
            "lingua_origine = CASE WHEN lingua_origine = 'manuale' "
            "  THEN lingua_origine ELSE ? END, "
            "gruppo_traduzione = ?, revisione = ?, "
            "data_documento = CASE WHEN data_origine = 'manuale' "
            "  THEN data_documento ELSE ? END, "
            "data_origine = CASE WHEN data_origine = 'manuale' "
            "  THEN data_origine ELSE ? END, "
            "ruolo_temporale = CASE WHEN ruolo_origine = 'manuale' "
            "  THEN ruolo_temporale ELSE ? END, "
            "ruolo_origine = CASE WHEN ruolo_origine = 'manuale' "
            "  THEN ruolo_origine ELSE ? END, "
            "data_aggiornamento = ? WHERE id = ?",
            (lingua, lingua_origine, gruppo, revisione, data, data_origine,
             ruolo, ruolo_origine, adesso, doc["id"]))
    conn.commit()
    return conteggi


def gruppi_traduzione_incompleti(conn):
    """Gruppi di traduzione con un solo documento. Non si accorpano per
    somiglianza: si segnalano, perche' la causa puo' essere una traduzione
    mancante oppure un nome che cambia da una lingua all'altra, e sono due
    problemi diversi."""
    return [dict(r) for r in conn.execute(
        "SELECT gruppo_traduzione, MIN(percorso_rel) AS esempio, "
        "COUNT(*) AS n FROM documenti WHERE stato = 'attivo' "
        "AND gruppo_traduzione <> '' GROUP BY gruppo_traduzione "
        "HAVING n = 1 ORDER BY gruppo_traduzione")]


# ===========================================================================
# 5. Il collegamento del documento alle entita' (molti a molti)
# ===========================================================================

def _famiglie_da_elenco(testo, chiavi):
    """Sigle ricavate da un elenco dichiarato nel documento, del tipo
    'Prodotti interessati: K230, S641 NH4'. Corrispondenza esatta, poi per
    prefisso univoco: se il prefisso vale per piu' famiglie l'esito e'
    ambiguo e non si collega nulla."""
    esiti = []
    per_norm = {_norm(k): k for k in chiavi}
    for pezzo in re.split(r"[,;]| e ", testo or ""):
        sigla = _norm(pezzo)
        if not sigla:
            continue
        if sigla in per_norm:
            esiti.append((per_norm[sigla], "esatto"))
            continue
        candidati = [k for n, k in per_norm.items() if n.startswith(sigla + " ")]
        if len(candidati) == 1:
            esiti.append((candidati[0], "prefisso"))
        elif len(candidati) > 1:
            esiti.append((None, "ambigua:%s" % sigla))
    return esiti


def _famiglie_da_nome(nome, chiavi):
    """Sigle riconosciute nel nome del file come sequenze di elementi
    consecutivi. Vince la sequenza piu' lunga, e gli elementi consumati non
    vengono riusati: cosi' 'K270-K280' produce due famiglie e non una."""
    per_norm = {_norm(k): k for k in chiavi}
    token = _token(nome)
    trovate, i = [], 0
    while i < len(token):
        preso = False
        for lunghezza in (3, 2, 1):
            if i + lunghezza > len(token):
                continue
            sigla = _norm(" ".join(token[i:i + lunghezza]))
            if sigla in per_norm:
                trovate.append(per_norm[sigla])
                i += lunghezza
                preso = True
                break
        if not preso:
            i += 1
    return trovate


def collega_documenti_entita(conn, radice):
    """Popola il collegamento molti a molti fra documenti ed entita'.

    Tre vie, in ordine di attendibilita' decrescente: il percorso di cartella
    per i documenti di prodotto, l'elenco dichiarato nel documento per le note
    tecniche, il nome del file come ultima risorsa. I collegamenti con origine
    'manuale' non vengono mai toccati."""
    tipo_famiglia = conn.execute(
        "SELECT id FROM tipi_entita WHERE nome = ?", (TIPO_FAMIGLIA,)).fetchone()
    if not tipo_famiglia:
        return {"documenti": 0}
    famiglie = {r["chiave"]: (r["id"], r["cartella_origine"])
                for r in conn.execute(
                    "SELECT id, chiave, cartella_origine FROM entita "
                    "WHERE tipo_id = ? AND stato = 'attiva'",
                    (tipo_famiglia[0],))}
    per_percorso = sorted(
        ((c, i) for i, c in famiglie.values() if c),
        key=lambda x: -len(x[0]))

    conn.execute("DELETE FROM documenti_entita WHERE origine <> 'manuale'")
    adesso = _adesso()
    conteggi = {"documenti": 0, "da_percorso": 0, "da_contenuto": 0,
                "da_nome": 0, "senza_entita": 0, "multipli": 0,
                "ambigui": 0}

    for doc in conn.execute("SELECT id, percorso_rel, ruolo_temporale "
                            "FROM documenti WHERE stato = 'attivo'").fetchall():
        conteggi["documenti"] += 1
        nome = os.path.basename(doc["percorso_rel"])
        collegate, origine = [], ""

        for cartella, eid in per_percorso:
            if doc["percorso_rel"].startswith(cartella + "/"):
                collegate, origine = [eid], "percorso"
                break

        if not collegate:
            righe = _testo_documento(radice, doc["percorso_rel"])
            elenco = _valore_etichetta(righe, ["Prodotti interessati"])
            if elenco:
                esiti = _famiglie_da_elenco(elenco, famiglie.keys())
                trovate = [famiglie[k][0] for k, e in esiti if k]
                conteggi["ambigui"] += sum(1 for k, e in esiti if not k)
                if trovate:
                    collegate, origine = trovate, "contenuto"

        if not collegate:
            trovate = [famiglie[k][0] for k in _famiglie_da_nome(
                nome, famiglie.keys())]
            if trovate:
                collegate, origine = trovate, "nome_file"

        if not collegate:
            conteggi["senza_entita"] += 1
            conn.execute("UPDATE documenti SET entita_id = NULL WHERE id = ?",
                         (doc["id"],))
            continue

        conteggi["da_%s" % origine] += 1
        if len(collegate) > 1:
            conteggi["multipli"] += 1
        for n, eid in enumerate(dict.fromkeys(collegate)):
            conn.execute(
                "INSERT OR IGNORE INTO documenti_entita (documento_id, "
                "entita_id, prevalente, origine, data) VALUES (?, ?, ?, ?, ?)",
                (doc["id"], eid, 1 if n == 0 else 0, origine, adesso))
        conn.execute("UPDATE documenti SET entita_id = ? WHERE id = ?",
                     (collegate[0], doc["id"]))
    conn.commit()
    return conteggi


# ===========================================================================
# 6. Il calcolo di vigenza
# ===========================================================================

def calcola_vigenza(conn):
    """Confronta le date e segnala i documenti candidati ad aggiornamento.

    Regola: per ogni entita', un documento descrittivo datato e' candidato ad
    aggiornamento se esiste, collegato alla stessa entita', un documento di
    ruolo 'aggiornamento' con data posteriore. Si registra la nota piu'
    recente fra quelle che lo superano, perche' una segnalazione per ogni nota
    sarebbe rumore.

    La segnalazione dice candidato ad aggiornamento, non dice obsoleto: e'
    un'ipotesi da sottoporre a chi presidia il documento. Per questo il
    ricalcolo rimuove le segnalazioni aperte non piu' prodotte, ma non tocca
    mai quelle su cui una persona si e' espressa."""
    adesso = _adesso()
    note, lingue_note = {}, {}
    for r in conn.execute(
            "SELECT de.entita_id AS eid, d.id AS did, d.data_documento AS data, "
            "d.lingua AS lingua "
            "FROM documenti d JOIN documenti_entita de ON de.documento_id = d.id "
            "WHERE d.stato = 'attivo' AND d.ruolo_temporale = ? "
            "AND d.data_documento <> ''", (RUOLO_AGGIORNAMENTO,)):
        note.setdefault(r["eid"], []).append((r["data"], r["did"]))
        lingue_note[r["did"]] = r["lingua"]

    prodotte = set()
    for r in conn.execute(
            "SELECT de.entita_id AS eid, d.id AS did, d.data_documento AS data, "
            "d.lingua AS lingua FROM documenti d "
            "JOIN documenti_entita de ON de.documento_id = d.id "
            "WHERE d.stato = 'attivo' AND d.ruolo_temporale = ? "
            "AND d.data_documento <> ''", (RUOLO_DESCRITTIVO,)).fetchall():
        posteriori = [n for n in note.get(r["eid"], []) if n[0] > r["data"]]
        if not posteriori:
            continue
        data_nota = max(n[0] for n in posteriori)
        pari_data = [n for n in posteriori if n[0] == data_nota]
        stessa_lingua = [n for n in pari_data
                         if lingue_note.get(n[1]) == r["lingua"]]
        causa = min(n[1] for n in (stessa_lingua or pari_data))
        scarto = (datetime.date.fromisoformat(data_nota)
                  - datetime.date.fromisoformat(r["data"])).days
        conn.execute(
            "INSERT OR IGNORE INTO segnalazioni_vigenza (entita_id, "
            "documento_id, causa_documento_id, lingua, motivo, scarto_giorni, "
            "data_calcolo, stato) VALUES (?, ?, ?, ?, 'nota_posteriore', ?, ?, "
            "'aperta')",
            (r["eid"], r["did"], causa, r["lingua"], scarto, adesso))
        conn.execute("UPDATE segnalazioni_vigenza SET scarto_giorni = ?, "
                     "data_calcolo = ? WHERE documento_id = ? AND "
                     "causa_documento_id = ? AND stato = 'aperta'",
                     (scarto, adesso, r["did"], causa))
        prodotte.add((r["did"], causa))

    for r in conn.execute("SELECT id, documento_id, causa_documento_id "
                          "FROM segnalazioni_vigenza WHERE stato = 'aperta'"
                          ).fetchall():
        if (r["documento_id"], r["causa_documento_id"]) not in prodotte:
            conn.execute("DELETE FROM segnalazioni_vigenza WHERE id = ?",
                         (r["id"],))

    conn.execute("UPDATE documenti SET stato_vigenza = 'non_valutato'")
    conn.execute(
        "UPDATE documenti SET stato_vigenza = 'vigente' "
        "WHERE ruolo_temporale = ? AND data_documento <> ''",
        (RUOLO_DESCRITTIVO,))
    conn.execute(
        "UPDATE documenti SET stato_vigenza = 'da_aggiornare' WHERE id IN "
        "(SELECT documento_id FROM segnalazioni_vigenza WHERE stato = 'aperta')")
    conn.commit()
    return misura_vigenza(conn)


def disponi_segnalazione(conn, seg_id, stato, nota="", utente=""):
    """Registra la disposizione umana su una segnalazione. Gli stati ammessi
    oltre ad 'aperta' sono 'confermata' (va aggiornata) e 'ignorata' (la nota
    non incide su questo documento). Sopravvivono al ricalcolo."""
    if stato not in ("aperta", "confermata", "ignorata"):
        return False
    conn.execute("UPDATE segnalazioni_vigenza SET stato = ?, nota = ?, "
                 "utente = ?, data_disposizione = ? WHERE id = ?",
                 (stato, nota, utente, _adesso(), seg_id))
    conn.commit()
    return True


def misura_vigenza(conn):
    """Gli indicatori del controllo: quanto e' calcolabile e che cosa dice."""
    tot = conn.execute("SELECT COUNT(*) FROM documenti WHERE stato = 'attivo'"
                       ).fetchone()[0]
    descrittivi = conn.execute(
        "SELECT COUNT(*) FROM documenti WHERE stato = 'attivo' "
        "AND ruolo_temporale = ?", (RUOLO_DESCRITTIVO,)).fetchone()[0]
    datati = conn.execute(
        "SELECT COUNT(*) FROM documenti WHERE stato = 'attivo' "
        "AND ruolo_temporale = ? AND data_documento <> ''",
        (RUOLO_DESCRITTIVO,)).fetchone()[0]
    aperte = [dict(r) for r in conn.execute(
        "SELECT s.id, s.scarto_giorni, s.lingua, e.chiave AS entita, "
        "d.percorso_rel AS documento, d.data_documento AS data_documento, "
        "c.percorso_rel AS causa, c.data_documento AS data_causa "
        "FROM segnalazioni_vigenza s "
        "JOIN documenti d ON d.id = s.documento_id "
        "JOIN documenti c ON c.id = s.causa_documento_id "
        "LEFT JOIN entita e ON e.id = s.entita_id "
        "WHERE s.stato = 'aperta' ORDER BY s.scarto_giorni DESC")]
    disposte = conn.execute(
        "SELECT COUNT(*) FROM segnalazioni_vigenza WHERE stato <> 'aperta'"
        ).fetchone()[0]
    scarti = sorted(s["scarto_giorni"] for s in aperte)
    return {
        "documenti": tot,
        "descrittivi": descrittivi,
        "descrittivi_datati": datati,
        "copertura_datazione": (0.0 if not descrittivi else
                                round(100.0 * datati / descrittivi, 1)),
        "segnalazioni_aperte": len(aperte),
        "segnalazioni_disposte": disposte,
        "scarto_mediano_giorni": (0 if not scarti
                                  else scarti[len(scarti) // 2]),
        "elenco": aperte,
    }


# ===========================================================================
# 7. Coerenza fra la data del file e quella dichiarata nel documento
# ===========================================================================
#
# La data di filesystem non e' la data del documento: usarla come tale
# produrrebbe un calcolo di vigenza fondato su un evento di sincronizzazione.
# Conservarla accanto a quella dichiarata, pero', e' utile: lo scarto fra le
# due e' un segnale. Un documento che si dichiara vecchio ma il cui file
# risulta toccato di recente puo' essere stato modificato senza aggiornare la
# data interna, riesportato, o semplicemente ricaricato. Non e' una regola,
# e' un indizio, e come tale va presentato.
#
# Il segnale ha senso solo dopo aver isolato le migrazioni massive: se un
# archivio e' stato trasferito in blocco, quasi tutti i file condividono la
# stessa giornata e lo scarto non distingue piu' nulla.

CHIAVE_QUOTA_MIGRAZIONE = "livello0_quota_migrazione"
CHIAVE_SOGLIA_SCARTO = "livello0_soglia_scarto_giorni"
QUOTA_MIGRAZIONE = 0.20      # quota di archivio in una sola giornata
MINIMO_MIGRAZIONE = 10       # sotto questo numero non e' una migrazione
SOGLIA_SCARTO = 180          # giorni oltre cui il ritocco merita un'occhiata

SEGNALE_MIGRAZIONE = "in_migrazione"
SEGNALE_COERENTE = "coerente"
SEGNALE_RITOCCATO = "ritoccato_dopo"
SEGNALE_ANTERIORE = "anteriore_al_contenuto"


def _impostazione_numero(conn, chiave, predefinito):
    grezzo = db.get_impostazione(conn, chiave, "")
    try:
        return type(predefinito)(grezzo) if grezzo.strip() else predefinito
    except (TypeError, ValueError):
        return predefinito


def giornate_di_migrazione(conn, quota=None):
    """Giornate che concentrano una quota anomala delle date di modifica.

    Su un archivio trasferito in blocco quasi tutti i file portano la stessa
    data: quella giornata non e' informativa per nessuno dei file che la
    condividono, e va isolata prima di leggere qualunque scarto.

    Oltre alla quota serve un minimo assoluto: su un archivio di pochi file
    la sola percentuale degenera, e due documenti emessi lo stesso giorno
    passerebbero per una migrazione. Meglio non riconoscerne una che
    sopprimere per errore i segnali di un archivio piccolo."""
    quota = QUOTA_MIGRAZIONE if quota is None else quota
    totale = conn.execute("SELECT COUNT(*) FROM inventario_file "
                          "WHERE stato = 'presente' AND data_modifica <> ''"
                          ).fetchone()[0]
    if not totale:
        return {}
    return {r["giorno"]: r["n"] for r in conn.execute(
        "SELECT substr(data_modifica, 1, 10) AS giorno, COUNT(*) AS n "
        "FROM inventario_file WHERE stato = 'presente' AND data_modifica <> '' "
        "GROUP BY giorno HAVING n >= ?",
        (max(MINIMO_MIGRAZIONE, int(totale * quota)),))}


def analizza_coerenza_date(conn):
    """Confronta, per ogni documento datato, la data dichiarata al suo interno
    con la data di modifica del file, e ne registra il segnale."""
    quota = _impostazione_numero(conn, CHIAVE_QUOTA_MIGRAZIONE,
                                 QUOTA_MIGRAZIONE)
    soglia = _impostazione_numero(conn, CHIAVE_SOGLIA_SCARTO, SOGLIA_SCARTO)
    giornate = giornate_di_migrazione(conn, quota)

    conteggi = {"in_migrazione": 0, "coerente": 0, "ritoccato_dopo": 0,
                "anteriore_al_contenuto": 0, "non_confrontabile": 0}
    for r in conn.execute(
            "SELECT d.id, d.data_documento, i.data_modifica "
            "FROM documenti d LEFT JOIN inventario_file i "
            "ON i.percorso_rel = d.percorso_rel "
            "WHERE d.stato = 'attivo'").fetchall():
        modifica = r["data_modifica"] or ""
        segnale, scarto = "", 0
        if modifica[:10] in giornate:
            segnale = SEGNALE_MIGRAZIONE
        elif not r["data_documento"] or not modifica:
            segnale = ""
        else:
            scarto = (datetime.date.fromisoformat(modifica[:10])
                      - datetime.date.fromisoformat(r["data_documento"])).days
            if scarto < 0:
                segnale = SEGNALE_ANTERIORE
            elif scarto > soglia:
                segnale = SEGNALE_RITOCCATO
            else:
                segnale = SEGNALE_COERENTE
        conteggi[segnale or "non_confrontabile"] += 1
        conn.execute("UPDATE documenti SET segnale_data = ?, "
                     "scarto_file_giorni = ? WHERE id = ?",
                     (segnale, scarto, r["id"]))
    conn.commit()
    conteggi["giornate_di_migrazione"] = giornate
    return conteggi


def misura_coerenza_date(conn):
    """Gli indizi da presentare, con l'elenco dei casi da guardare."""
    quota = _impostazione_numero(conn, CHIAVE_QUOTA_MIGRAZIONE,
                                 QUOTA_MIGRAZIONE)
    giornate = giornate_di_migrazione(conn, quota)
    conteggi = {r["segnale_data"] or "non_confrontabile": r["n"]
                for r in conn.execute(
                    "SELECT segnale_data, COUNT(*) AS n FROM documenti "
                    "WHERE stato = 'attivo' GROUP BY segnale_data")}
    casi = [dict(r) for r in conn.execute(
        "SELECT d.percorso_rel, d.data_documento, d.data_origine, "
        "d.segnale_data, d.scarto_file_giorni, i.data_modifica "
        "FROM documenti d LEFT JOIN inventario_file i "
        "ON i.percorso_rel = d.percorso_rel "
        "WHERE d.stato = 'attivo' AND d.segnale_data IN (?, ?) "
        "ORDER BY d.scarto_file_giorni DESC",
        (SEGNALE_RITOCCATO, SEGNALE_ANTERIORE))]
    totale = conn.execute("SELECT COUNT(*) FROM documenti WHERE stato = "
                          "'attivo'").fetchone()[0]
    in_migrazione = conteggi.get(SEGNALE_MIGRAZIONE, 0)
    return {
        "totale": totale,
        "giornate_di_migrazione": sorted(giornate.items()),
        "in_migrazione": in_migrazione,
        "quota_in_migrazione": (0.0 if not totale else
                                round(100.0 * in_migrazione / totale, 1)),
        "coerenti": conteggi.get(SEGNALE_COERENTE, 0),
        "ritoccati": conteggi.get(SEGNALE_RITOCCATO, 0),
        "anteriori": conteggi.get(SEGNALE_ANTERIORE, 0),
        "non_confrontabili": conteggi.get("non_confrontabile", 0),
        "casi": casi,
    }
